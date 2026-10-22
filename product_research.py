#!/usr/bin/env python3
"""户外露营轻量化储能 · 公开数据选品自动化模块。

边界声明
--------
1. 仅处理 Amazon / AliExpress / 1688 的公开可见搜索页或本地公开快照；
2. 不登录、不使用卖家后台、不调用任何店铺 API；
3. 不绕过验证码、不伪造 Cookie、不使用代理池对抗访问控制；
4. live 抓取前检查 robots.txt，失败时明确记录并停止，不静默绕过；
5. 演示默认使用 ``data/mock/public_research_snapshot.json``，数据均标注为模拟。

用法
----
    python product_research.py
    python product_research.py --mode offline --min-margin 30 --top-n 20
    python product_research.py --mode live --limit 20
    python product_research.py --help
"""

from __future__ import annotations

import argparse
import copy
import json
import math
import re
import statistics
import sys
import time
from collections import Counter, defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus, urljoin, urlsplit
from urllib.robotparser import RobotFileParser

# ---------------------------------------------------------------------------
# 路径与可选依赖
# ---------------------------------------------------------------------------
PROJECT_ROOT = Path(__file__).resolve().parent
VENDOR_DIR = PROJECT_ROOT / ".vendor"
if VENDOR_DIR.is_dir() and str(VENDOR_DIR) not in sys.path:
    sys.path.insert(0, str(VENDOR_DIR))

import utils  # noqa: E402

logger = utils.get_logger("product_research")

try:  # Excel 是硬依赖，但延迟到导出时给出更清晰的错误
    import openpyxl  # noqa: F401
except ImportError:  # pragma: no cover - 运行环境缺依赖
    openpyxl = None  # type: ignore[assignment]

DEFAULT_CONFIG = PROJECT_ROOT / "config" / "product_research.yaml"
PLATFORM_ALIASES = {
    "amazon": "amazon",
    "amz": "amazon",
    "aliexpress": "aliexpress",
    "ali": "aliexpress",
    "1688": "alibaba_1688",
    "alibaba_1688": "alibaba_1688",
}
SKU_PREFIX = {"amazon": "AMZ", "aliexpress": "ALI", "alibaba_1688": "1688"}
CAPACITY_BUCKETS = ((100, 199), (200, 299), (300, 399), (400, 500))


# ---------------------------------------------------------------------------
# 基础类型与配置
# ---------------------------------------------------------------------------
class ResearchError(RuntimeError):
    """选品模块业务异常。"""


class CrawlError(ResearchError):
    """公开页面抓取失败。"""


class RobotsDenied(CrawlError):
    """robots.txt 不允许抓取。"""


class ParseError(ResearchError):
    """公开页面结构变化，未能解析出商品。"""


@dataclass(frozen=True)
class ResearchConfig:
    root: Path
    path: Path
    data: dict[str, Any]

    def get(self, dotted: str, default: Any = None) -> Any:
        return utils.deep_get(self.data, dotted, default)

    def section(self, name: str) -> dict[str, Any]:
        value = self.get(name, {})
        return dict(value) if isinstance(value, dict) else {}

    def platform_items(self) -> list[tuple[str, dict[str, Any]]]:
        platforms = self.section("platforms")
        return [
            (str(key), dict(value))
            for key, value in platforms.items()
            if isinstance(value, dict) and value.get("enabled", True)
        ]

    def resolve(self, dotted: str) -> Path:
        raw = self.get(dotted)
        if not raw:
            raise ResearchError(f"缺少路径配置: {dotted}")
        path = Path(str(raw)).expanduser()
        return path if path.is_absolute() else (self.root / path).resolve()


def load_research_config(
    config_path: str | Path | None = None,
    *,
    root: str | Path | None = None,
) -> ResearchConfig:
    project_root = Path(root).expanduser().resolve() if root else PROJECT_ROOT
    path = Path(config_path).expanduser() if config_path else project_root / "config" / "product_research.yaml"
    if not path.is_absolute():
        path = (project_root / path).resolve()
    data = utils.read_yaml(path)
    if not isinstance(data, dict):
        raise ResearchError(f"配置文件根节点必须是映射: {path}")
    if not data.get("platforms"):
        raise ResearchError("配置缺少 platforms")
    if not data.get("filters"):
        raise ResearchError("配置缺少 filters")
    return ResearchConfig(root=project_root, path=path, data=data)


def _num(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", str(value))
    if not match:
        return default
    try:
        return float(match.group(0).replace(",", ""))
    except ValueError:
        return default


def _int(value: Any, default: int = 0) -> int:
    return int(round(_num(value, float(default))))


def _clamp(value: float, low: float = 0.0, high: float = 100.0) -> float:
    return max(low, min(high, value))


def _bucket(capacity: float) -> str:
    for low, high in CAPACITY_BUCKETS:
        if low <= capacity <= high:
            return f"{low}-{high}"
    return "out_of_range"


def _title_key(title: str) -> str:
    return re.sub(r"[^\w]+", "", str(title).lower(), flags=re.UNICODE)


def _extract_spec(title: str, field: str) -> float:
    """从公开标题中解析容量/重量/功率，作为字段缺失时的回退。"""

    text = str(title)
    if field == "capacity_wh":
        match = re.search(r"(\d{2,4}(?:\.\d+)?)\s*wh\b", text, flags=re.IGNORECASE)
        if match:
            return _num(match.group(1))
        match = re.search(r"(\d+(?:\.\d+)?)\s*kwh\b", text, flags=re.IGNORECASE)
        return _num(match.group(1)) * 1000 if match else 0.0
    if field == "weight_kg":
        match = re.search(r"(\d+(?:\.\d+)?)\s*kg\b", text, flags=re.IGNORECASE)
        return _num(match.group(1)) if match else 0.0
    if field == "output_w":
        match = re.search(r"(\d{2,4})\s*w\b", text, flags=re.IGNORECASE)
        return _num(match.group(1)) if match else 0.0
    return 0.0


def _parse_rating(text: str) -> float:
    match = re.search(r"([0-5](?:\.\d)?)\s*(?:out of 5|/5|分|星)?", str(text))
    return _num(match.group(1)) if match else 0.0


# ---------------------------------------------------------------------------
# 清洗
# ---------------------------------------------------------------------------
class ProductCleaner:
    """把公开页面/快照记录清洗成统一分析表。"""

    def __init__(self, config: ResearchConfig) -> None:
        self.config = config
        self.filters = config.section("filters")
        self.platforms = {key: value for key, value in config.platform_items()}
        self.fx = config.get("cost_model.fx_to_usd", {}) or {}

    def clean(self, raw_records: Iterable[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        kept: dict[tuple[str, str], dict[str, Any]] = {}
        drops: Counter[str] = Counter()
        raw_by_platform: Counter[str] = Counter()
        clean_by_platform: Counter[str] = Counter()
        seen_ids: set[tuple[str, str]] = set()
        seen_titles: set[tuple[str, str]] = set()

        for raw in raw_records:
            if not isinstance(raw, dict):
                drops["invalid_record"] += 1
                continue
            platform = PLATFORM_ALIASES.get(str(raw.get("platform", "")).lower(), "")
            if not platform or platform not in self.platforms:
                drops["unknown_platform"] += 1
                continue
            raw_by_platform[platform] += 1
            title = re.sub(r"\s+", " ", str(raw.get("title", ""))).strip()
            if not title:
                drops["missing_title"] += 1
                continue

            lower_title = title.lower()
            excluded = next(
                (word for word in self.filters.get("exclude_keywords", []) if str(word).lower() in lower_title),
                None,
            )
            if excluded:
                drops["excluded_keyword"] += 1
                continue

            price_local = _num(raw.get("price_local", raw.get("price", 0)))
            if price_local <= 0:
                drops["missing_price"] += 1
                continue

            platform_cfg = self.platforms[platform]
            currency = str(raw.get("currency") or platform_cfg.get("currency", "USD")).upper()
            fx = _num(self.fx.get(currency), 1.0)
            price_usd = price_local * fx

            capacity = _num(raw.get("capacity_wh")) or _extract_spec(title, "capacity_wh")
            if capacity <= 0:
                drops["missing_capacity"] += 1
                continue
            output_w = _num(raw.get("output_w")) or _extract_spec(title, "output_w")
            weight = _num(raw.get("weight_kg")) or _extract_spec(title, "weight_kg")
            weight_source = "public_field" if _num(raw.get("weight_kg")) else "title_parse"
            if weight <= 0:
                weight = 5.0
                weight_source = "category_default"

            low_cap, high_cap = [float(x) for x in self.filters.get("capacity_wh", [100, 500])]
            if not low_cap <= capacity <= high_cap:
                drops["capacity_out_of_range"] += 1
                continue
            max_weight = float(self.filters.get("max_weight_kg", 15))
            if weight > max_weight:
                drops["weight_out_of_range"] += 1
                continue
            low_price, high_price = [float(x) for x in self.filters.get("price_usd", [49, 1500])]
            if not low_price <= price_usd <= high_price:
                drops["price_out_of_range"] += 1
                continue
            price_per_wh = price_usd / capacity
            low_ppw, high_ppw = [float(x) for x in self.filters.get("price_per_wh_usd", [0.1, 8])]
            if not low_ppw <= price_per_wh <= high_ppw:
                drops["price_per_wh_outlier"] += 1
                continue
            rating = _num(raw.get("rating"))
            if rating and rating < float(self.filters.get("min_rating", 0)):
                drops["rating_below_threshold"] += 1
                continue

            product_id = str(raw.get("product_id") or raw.get("sku") or raw.get("record_id") or "")
            if not product_id:
                product_id = utils.safe_filename(_title_key(title))[:40]
            identity = (platform, product_id)
            if identity in seen_ids:
                drops["duplicate_product_id"] += 1
                continue

            key = (platform, str(raw.get("product_id") or _title_key(title)))
            title_identity = (platform, _title_key(title))
            if title_identity in seen_titles:
                drops["duplicate_title"] += 1
                continue
            missing = [
                name
                for name, value in (
                    ("url", raw.get("url")),
                    ("monthly_search_volume", raw.get("monthly_search_volume")),
                    ("competitor_count", raw.get("competitor_count")),
                    ("review_count", raw.get("review_count")),
                )
                if not value
            ]
            quality = 100.0 - 8.0 * len(missing)
            if weight_source == "category_default":
                quality -= 12
            if not raw.get("brand"):
                quality -= 3
            quality = _clamp(quality)

            row = {
                "record_id": str(raw.get("record_id") or product_id),
                "platform": platform,
                "platform_label": str(platform_cfg.get("label", platform)),
                "product_id": product_id,
                "sku": f"{SKU_PREFIX.get(platform, platform.upper()[:3])}-{product_id}",
                "title": title,
                "brand": str(raw.get("brand") or "Unknown"),
                "url": str(raw.get("url") or raw.get("source_page") or ""),
                "source_page": str(raw.get("source_page") or raw.get("url") or ""),
                "market": str(raw.get("market") or platform_cfg.get("market", "US")).upper(),
                "currency": currency,
                "price_local": round(price_local, 2),
                "price_usd": round(price_usd, 2),
                "capacity_wh": round(capacity, 1),
                "output_w": round(output_w, 1),
                "weight_kg": round(weight, 2),
                "weight_source": weight_source,
                "battery_type": str(raw.get("battery_type") or "Unknown"),
                "solar_input_w": round(_num(raw.get("solar_input_w")), 1),
                "ports": str(raw.get("ports") or ""),
                "rating": round(rating, 2),
                "review_count": _int(raw.get("review_count")),
                "seller": str(raw.get("seller") or ""),
                "monthly_search_volume": _int(raw.get("monthly_search_volume")),
                "competitor_count": _int(raw.get("competitor_count")),
                "ad_count": _int(raw.get("ad_count")),
                "moq": max(1, _int(raw.get("moq"), 1)),
                "is_sourcing": bool(platform_cfg.get("is_sourcing", False)),
                "fee_rate": _num(platform_cfg.get("fee_rate"), 0.15),
                "benchmark_cost_ratio": _num(platform_cfg.get("benchmark_cost_ratio"), 0.42),
                "conversion_rate": _num(platform_cfg.get("conversion_rate"), 0.02),
                "price_per_wh_usd": round(price_per_wh, 4),
                "price_band": "low" if price_usd < 150 else ("mid" if price_usd <= 400 else "high"),
                "demand_source": "public_search_proxy" if raw.get("monthly_search_volume") else "missing",
                "data_quality_score": round(quality, 2),
                "missing_fields": ",".join(missing),
                "data_source": str(raw.get("data_source") or "offline_public_snapshot"),
                "captured_at": str(raw.get("collected_at") or ""),
            }

            previous = kept.get(key)
            if previous is not None:
                drops["duplicate_title"] += 1
                if row["review_count"] > previous["review_count"]:
                    kept[key] = row
                continue
            seen_ids.add(identity)
            seen_titles.add(title_identity)
            kept[key] = row

        rows = list(kept.values())
        rows.sort(key=lambda item: (item["platform"], -item["review_count"]))
        for row in rows:
            clean_by_platform[row["platform"]] += 1
        stats = {
            "raw_records": sum(raw_by_platform.values()),
            "clean_records": len(rows),
            "dropped_records": sum(drops.values()),
            "dropped_by_reason": dict(drops.most_common()),
            "by_platform_raw": dict(raw_by_platform),
            "by_platform_clean": dict(clean_by_platform),
        }
        return rows, stats


# ---------------------------------------------------------------------------
# 指标计算
# ---------------------------------------------------------------------------
class SelectionMetrics:
    """计算单 SKU 毛利、市场容量、竞争度与综合机会分。

    所有指标都是可解释的规则模型，不依赖平台私有 API。关键公式同时写入
    Excel 的“指标说明”页，便于业务复核。
    """

    def __init__(self, config: ResearchConfig) -> None:
        self.config = config
        self.cost = config.section("cost_model")
        self.competition_cfg = config.section("competition")
        self.capacity_cfg = config.section("market_capacity")
        self.opportunity_cfg = config.section("opportunity")
        self.thresholds = dict(self.opportunity_cfg.get("thresholds", {}) or {})

    def enrich(self, rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
        if not rows:
            raise ResearchError("清洗后没有可用商品，无法计算选品指标")
        retail_rows = [row for row in rows if not row["is_sourcing"]]
        benchmark = self._retail_benchmarks(retail_rows or rows)
        cv_by_bucket, cv_by_platform = self._price_dispersion(retail_rows or rows)
        total_weight = sum(float(v) for v in (self.opportunity_cfg.get("weights") or {}).values()) or 1.0

        enriched: list[dict[str, Any]] = []
        for source in rows:
            row = dict(source)
            bucket = _bucket(row["capacity_wh"])
            is_sourcing = bool(row["is_sourcing"])
            retail_price = benchmark["by_bucket"].get(bucket, benchmark["overall"])
            benchmark_source = "capacity_bucket_median" if bucket in benchmark["by_bucket"] else "overall_median"
            if not is_sourcing:
                retail_price = row["price_usd"]
                benchmark_source = "own_listing_price"
            retail_price = round(float(retail_price), 2)

            supply_cost = row["price_usd"] if is_sourcing else row["price_usd"] * row["benchmark_cost_ratio"]
            supply_cost = round(float(supply_cost), 2)
            weight = float(row["weight_kg"])
            freight = round(
                max(
                    float(self.cost.get("minimum_freight_usd", 6.0)),
                    weight * float(self.cost.get("inbound_freight_per_kg_usd", 2.2)),
                ),
                2,
            )
            duty_rate = float((self.cost.get("duty_pct_by_market") or {}).get(row["market"], 3.0))
            duty = round(supply_cost * duty_rate / 100.0, 2)
            packaging = round(float(self.cost.get("packaging_usd", 1.5)), 2)
            landed = round(supply_cost + freight + duty + packaging, 2)
            fee = round(retail_price * float(row["fee_rate"]), 2)
            fulfillment = round(
                float(self.cost.get("fulfillment_base_usd", 3.0))
                + weight * float(self.cost.get("fulfillment_per_kg_usd", 1.8)),
                2,
            )
            gross_profit = round(retail_price - landed - fee - fulfillment, 2)
            gross_margin = round(gross_profit / retail_price * 100.0, 2) if retail_price else 0.0
            ad_cost = round(retail_price * float(self.cost.get("ad_cost_ratio", 0.08)), 2)
            contribution = round(gross_profit - ad_cost, 2)
            contribution_margin = round(contribution / retail_price * 100.0, 2) if retail_price else 0.0

            search_volume = int(row["monthly_search_volume"])
            conversion = float(row["conversion_rate"])
            monthly_units = round(search_volume * conversion, 0)
            market_capacity = round(monthly_units * retail_price, 2)
            addressable_share = float(self.capacity_cfg.get("addressable_share", 0.15))
            addressable_units = round(monthly_units * addressable_share, 0)
            addressable_revenue = round(addressable_units * retail_price, 2)

            price_cv = cv_by_bucket.get(bucket, 0.0) if is_sourcing else cv_by_platform.get(
                (row["platform"], bucket), cv_by_bucket.get(bucket, 0.0)
            )
            competition = self._competition_score(row, price_cv)
            margin_score = _clamp(gross_margin / 45.0 * 100.0)
            market_score = _clamp(market_capacity / 500_000.0 * 100.0)
            competition_opportunity = _clamp(100.0 - competition)
            lightweight_score = _clamp(100.0 - max(0.0, weight - 2.0) * 7.5)
            data_quality = _clamp(float(row["data_quality_score"]))
            weights = self.opportunity_cfg.get("weights", {}) or {}
            opportunity = (
                margin_score * float(weights.get("margin", 0.35))
                + market_score * float(weights.get("market", 0.25))
                + competition_opportunity * float(weights.get("competition", 0.20))
                + lightweight_score * float(weights.get("lightweight", 0.10))
                + data_quality * float(weights.get("data_quality", 0.10))
            ) / total_weight

            gates: list[str] = []
            if gross_margin < float(self.thresholds.get("min_gross_margin_pct", 25)):
                gates.append(f"毛利率<{self.thresholds.get('min_gross_margin_pct', 25)}%")
            if market_capacity < float(self.thresholds.get("min_market_capacity_usd", 120000)):
                gates.append("市场容量不足")
            if competition > float(self.thresholds.get("max_competition_score", 70)):
                gates.append("竞争度过高")
            if weight > float(self.thresholds.get("max_weight_kg", 12)):
                gates.append("重量超限")
            if opportunity < float(self.thresholds.get("min_opportunity_score", 65)):
                gates.append("综合分不足")

            row.update(
                {
                    "retail_price_usd": retail_price,
                    "retail_benchmark_source": benchmark_source,
                    "supply_cost_usd": supply_cost,
                    "inbound_freight_usd": freight,
                    "duty_usd": duty,
                    "packaging_usd": packaging,
                    "landed_cost_usd": landed,
                    "platform_fee_usd": fee,
                    "fulfillment_cost_usd": fulfillment,
                    "gross_profit_usd": gross_profit,
                    "gross_margin_pct": gross_margin,
                    "ad_cost_est_usd": ad_cost,
                    "contribution_profit_usd": contribution,
                    "contribution_margin_pct": contribution_margin,
                    "monthly_units": int(monthly_units),
                    "market_capacity_usd": market_capacity,
                    "addressable_units": int(addressable_units),
                    "addressable_revenue_usd": addressable_revenue,
                    "price_dispersion_pct": round(price_cv, 2),
                    "competition_score": round(competition, 2),
                    "margin_score": round(margin_score, 2),
                    "market_score": round(market_score, 2),
                    "competition_opportunity_score": round(competition_opportunity, 2),
                    "lightweight_score": round(lightweight_score, 2),
                    "opportunity_score": round(opportunity, 2),
                    "gate_failures": gates,
                    "gate_status": "通过" if not gates else "未通过：" + "；".join(gates),
                    "high_potential": not gates,
                    "recommendation": self._recommendation(opportunity, gates),
                    "suggested_action": self._suggested_action(row, is_sourcing, gates),
                    "compliance_note": self._compliance_note(row["market"]),
                    "cost_basis": (
                        "1688 公开批发价 + 目标市场零售中位价"
                        if is_sourcing
                        else f"{row['platform_label']} 公开零售价 × 成本率 {row['benchmark_cost_ratio']:.0%}"
                    ),
                    "target_channel": (
                        "Amazon US（1688 货源假设）" if is_sourcing else f"{row['platform_label']} {row['market']}"
                    ),
                }
            )
            enriched.append(row)

        enriched.sort(key=lambda item: (-item["opportunity_score"], item["competition_score"]))
        for index, row in enumerate(enriched, start=1):
            row["rank"] = index
        high = [row for row in enriched if row["high_potential"]][: int(self.opportunity_cfg.get("top_n", 30))]
        passing_count = sum(1 for row in enriched if row["high_potential"])
        gate_counts = Counter(reason.split("：")[0] for row in enriched for reason in row["gate_failures"])
        stats = {
            "total": len(enriched),
            "high_potential": len(high),
            "high_potential_total": passing_count,
            "gate_failures": dict(gate_counts),
            "avg_margin_pct": round(statistics.fmean(row["gross_margin_pct"] for row in enriched), 2),
            "avg_competition": round(statistics.fmean(row["competition_score"] for row in enriched), 2),
            "median_market_capacity": round(statistics.median(row["market_capacity_usd"] for row in enriched), 2),
            "benchmark": {key: round(value, 2) for key, value in benchmark["by_bucket"].items()},
        }
        return enriched, high, stats

    def _retail_benchmarks(self, rows: list[dict[str, Any]]) -> dict[str, Any]:
        by_bucket: dict[str, list[float]] = defaultdict(list)
        for row in rows:
            by_bucket[_bucket(row["capacity_wh"])].append(float(row["price_usd"]))
        medians = {key: statistics.median(values) for key, values in by_bucket.items() if values}
        all_prices = [float(row["price_usd"]) for row in rows if row["price_usd"] > 0]
        return {
            "by_bucket": medians,
            "overall": statistics.median(all_prices) if all_prices else 0.0,
        }

    @staticmethod
    def _price_dispersion(rows: list[dict[str, Any]]) -> tuple[dict[str, float], dict[tuple[str, str], float]]:
        by_bucket: dict[str, list[float]] = defaultdict(list)
        by_platform: dict[tuple[str, str], list[float]] = defaultdict(list)
        for row in rows:
            price = float(row["price_usd"])
            bucket = _bucket(row["capacity_wh"])
            by_bucket[bucket].append(price)
            by_platform[(row["platform"], bucket)].append(price)

        def cv(values: list[float]) -> float:
            if len(values) < 2:
                return 0.0
            mean = statistics.fmean(values)
            return statistics.pstdev(values) / mean * 100.0 if mean else 0.0

        return (
            {key: cv(values) for key, values in by_bucket.items()},
            {key: cv(values) for key, values in by_platform.items()},
        )

    def _competition_score(self, row: dict[str, Any], price_cv: float) -> float:
        weights = self.competition_cfg.get("weights", {}) or {}
        density = _clamp(
            math.log1p(max(0, int(row["competitor_count"])))
            / math.log1p(float(self.competition_cfg.get("saturation_competitors", 1000)))
            * 100.0
        )
        review = _clamp(
            math.log1p(max(0, int(row["review_count"])))
            / math.log1p(float(self.competition_cfg.get("review_ceiling", 50000)))
            * 100.0
        )
        ad = _clamp(
            int(row["ad_count"]) / max(1.0, float(self.competition_cfg.get("ad_ceiling", 60))) * 100.0
        )
        price = _clamp(price_cv / max(1.0, float(self.competition_cfg.get("price_cv_full_scale_pct", 60))) * 100.0)
        total = sum(float(v) for v in weights.values()) or 1.0
        return (
            density * float(weights.get("listing_density", 0.35))
            + review * float(weights.get("review_barrier", 0.30))
            + ad * float(weights.get("ad_pressure", 0.20))
            + price * float(weights.get("price_dispersion", 0.15))
        ) / total

    @staticmethod
    def _recommendation(score: float, gates: list[str]) -> str:
        if gates:
            return "未达门槛"
        if score >= 80:
            return "S-优先打样"
        if score >= 70:
            return "A-小批量测试"
        return "B-进入观察池"

    @staticmethod
    def _suggested_action(row: dict[str, Any], is_sourcing: bool, gates: list[str]) -> str:
        if gates:
            return "先补齐：" + "；".join(gates)
        if is_sourcing:
            return "联系 1688 供应商索样，核验容量/电芯/BMS/认证，再按目标市场零售价复核"
        return "对标公开零售页，补充供应商报价、头程与退货率后复核毛利"

    @staticmethod
    def _compliance_note(market: str) -> str:
        mapping = {
            "US": "需核验 UL 2743 / FCC / UN38.3 / Prop 65；本表不构成认证结论",
            "DE": "需核验 CE / RoHS / WEEE / UN38.3 / EU Battery Regulation；本表不构成认证结论",
            "FR": "需核验 CE / RoHS / WEEE / UN38.3 / EU Battery Regulation；本表不构成认证结论",
            "UK": "需核验 UKCA / RoHS / WEEE / UN38.3；本表不构成认证结论",
        }
        return mapping.get(market, "需按目标市场核验电池、电气与运输合规；本表不构成认证结论")


# ---------------------------------------------------------------------------
# 公开页面抓取：robots.txt + 限速 + 不绕过访问控制
# ---------------------------------------------------------------------------
class PublicFetcher:
    """只读取公开 HTML，主动遵守 robots.txt 与请求间隔。"""

    def __init__(self, config: ResearchConfig) -> None:
        crawler = config.section("crawler")
        self.user_agent = str(
            crawler.get("user_agent", "OutdoorEssAI/0.2 (public-research; contact: research@example.com)")
        )
        self.timeout = float(crawler.get("timeout_seconds", 20))
        self.delay = max(0.0, float(crawler.get("delay_seconds", 2.0)))
        self.respect_robots = bool(crawler.get("respect_robots", True))
        self._robots: dict[str, bool] = {}
        self._last_request: dict[str, float] = {}
        self._session: Any = None
        self._requests: Any = None

    def _get_session(self) -> Any:
        if self._requests is None:
            try:
                import requests as _requests
            except ImportError as exc:  # pragma: no cover
                raise CrawlError("live 模式需要 requests：pip install requests") from exc
            self._requests = _requests
            self._session = _requests.Session()
            self._session.headers.update(
                {
                    "User-Agent": self.user_agent,
                    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                    "Accept-Language": "en-US,en;q=0.9,zh-CN;q=0.8",
                    "Cache-Control": "no-cache",
                }
            )
        return self._session

    def _robots_allowed(self, url: str) -> bool:
        if not self.respect_robots:
            return True
        parts = urlsplit(url)
        origin = f"{parts.scheme}://{parts.netloc}"
        if origin in self._robots:
            return self._robots[origin]
        robots_url = f"{origin}/robots.txt"
        try:
            response = self._get_session().get(robots_url, timeout=self.timeout)
        except Exception as exc:
            raise CrawlError(f"读取 robots.txt 失败：{robots_url}：{exc}") from exc
        if response.status_code == 404:
            allowed = True
        elif response.status_code in {401, 403}:
            allowed = False
        elif not 200 <= response.status_code < 300:
            allowed = False
        else:
            parser = RobotFileParser()
            parser.set_url(robots_url)
            parser.parse(response.text.splitlines())
            allowed = parser.can_fetch(self.user_agent, url)
        self._robots[origin] = allowed
        return allowed

    def get(self, url: str) -> str:
        if not self._robots_allowed(url):
            raise RobotsDenied(
                f"robots.txt 不允许抓取 {url}；已停止，未尝试绕过。"
            )
        parts = urlsplit(url)
        origin = f"{parts.scheme}://{parts.netloc}"
        last = self._last_request.get(origin, 0.0)
        wait = self.delay - (time.monotonic() - last)
        if wait > 0:
            time.sleep(wait)
        try:
            response = self._get_session().get(url, timeout=self.timeout, allow_redirects=True)
        except Exception as exc:
            raise CrawlError(f"公开页面请求异常：{url}：{exc}") from exc
        self._last_request[origin] = time.monotonic()
        if response.status_code in {401, 403, 429, 503}:
            raise CrawlError(
                f"公开页面拒绝访问（HTTP {response.status_code}）：{url}；"
                "不进行验证码、Cookie 或代理绕过。"
            )
        if response.status_code >= 400:
            raise CrawlError(f"公开页面请求失败（HTTP {response.status_code}）：{url}")
        text = response.text or ""
        lowered = text[:5000].lower()
        if "captcha" in lowered or "robot check" in lowered or "verify you are human" in lowered:
            raise CrawlError(f"公开页面返回验证码/人机校验：{url}；已停止，不做绕过。")
        return text

    def close(self) -> None:
        if self._session is not None:
            self._session.close()
            self._session = None


def _soup(html: str) -> Any:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover
        raise CrawlError("live 模式需要 beautifulsoup4：pip install beautifulsoup4") from exc
    return BeautifulSoup(html, "html.parser")


def _walk_jsonld(value: Any) -> Iterable[dict[str, Any]]:
    if isinstance(value, dict):
        types = value.get("@type")
        if types == "Product" or (isinstance(types, list) and "Product" in types):
            yield value
        for child in value.values():
            yield from _walk_jsonld(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_jsonld(child)


def _jsonld_products(html: str, page_url: str) -> list[dict[str, Any]]:
    soup = _soup(html)
    products: list[dict[str, Any]] = []
    for script in soup.find_all("script", attrs={"type": "application/ld+json"}):
        try:
            payload = json.loads(script.string or script.get_text() or "")
        except (TypeError, ValueError):
            continue
        for item in _walk_jsonld(payload):
            offers = item.get("offers")
            if isinstance(offers, list):
                offers = offers[0] if offers else {}
            if not isinstance(offers, dict):
                offers = {}
            price = _num(offers.get("price", offers.get("lowPrice", 0)))
            if price <= 0:
                continue
            rating = item.get("aggregateRating") or {}
            url = urljoin(page_url, str(item.get("url") or page_url))
            brand = item.get("brand")
            if isinstance(brand, dict):
                brand = brand.get("name")
            products.append(
                {
                    "title": str(item.get("name") or "").strip(),
                    "brand": str(brand or ""),
                    "url": url,
                    "price_local": price,
                    "currency": str(offers.get("priceCurrency") or "USD").upper(),
                    "rating": _num(rating.get("ratingValue")),
                    "review_count": _int(rating.get("reviewCount") or rating.get("ratingCount")),
                    "seller": str((offers.get("seller") or {}).get("name") or "")
                    if isinstance(offers.get("seller"), dict)
                    else "",
                }
            )
    return [item for item in products if item["title"] and item["price_local"] > 0]


class BasePlatformAdapter:
    key = ""
    label = ""

    def __init__(self, key: str, platform_cfg: dict[str, Any], fetcher: PublicFetcher) -> None:
        self.key = key
        self.cfg = platform_cfg
        self.label = str(platform_cfg.get("label", key))
        self.fetcher = fetcher

    def search_url(self, query: str | None = None) -> str:
        query = query or str(self.cfg.get("query", "portable power station"))
        template = str(self.cfg.get("search_url", ""))
        return template.format(query=quote_plus(query))

    def crawl(self, *, limit: int, query: str | None = None) -> list[dict[str, Any]]:
        page_url = self.search_url(query)
        html = self.fetcher.get(page_url)
        records = self.parse(html, page_url)
        if not records:
            raise ParseError(
                f"{self.label} 公开搜索页未解析出商品，可能页面结构变化或被要求登录。"
            )
        for record in records:
            record.setdefault("platform", self.key)
            record.setdefault("market", self.cfg.get("market", "US"))
            record.setdefault("currency", self.cfg.get("currency", "USD"))
            record.setdefault("source_page", page_url)
            record.setdefault("data_source", "public_live_page")
            record.setdefault("collected_at", utils.utc_now_iso())
            record.setdefault("competitor_count", len(records))
        return records[:limit]

    def parse(self, html: str, page_url: str) -> list[dict[str, Any]]:
        raise NotImplementedError

    def _record_from_jsonld(self, item: dict[str, Any], page_url: str) -> dict[str, Any]:
        return {
            "platform": self.key,
            "product_id": self._product_id(item.get("url", ""), item.get("title", "")),
            "title": item.get("title", ""),
            "brand": item.get("brand", ""),
            "url": item.get("url") or page_url,
            "market": self.cfg.get("market", "US"),
            "currency": item.get("currency") or self.cfg.get("currency", "USD"),
            "price_local": item.get("price_local", 0),
            "rating": item.get("rating", 0),
            "review_count": item.get("review_count", 0),
            "seller": item.get("seller", ""),
            "monthly_search_volume": 0,
            "competitor_count": 0,
            "ad_count": 0,
            "source_page": page_url,
            "data_source": "public_live_page",
            "collected_at": utils.utc_now_iso(),
        }

    @staticmethod
    def _product_id(url: str, title: str) -> str:
        match = re.search(r"(?:/dp/|/item/|offer/)([A-Za-z0-9_-]{5,})", str(url))
        if match:
            return match.group(1)
        return utils.safe_filename(_title_key(title))[:32] or "unknown"


class AmazonAdapter(BasePlatformAdapter):
    def parse(self, html: str, page_url: str) -> list[dict[str, Any]]:
        soup = _soup(html)
        page_text = soup.get_text(" ", strip=True)
        count_match = re.search(r"(?:over|about)\s+([\d,]+)\s+results", page_text, re.IGNORECASE)
        page_count = _int(count_match.group(1)) if count_match else 0
        records: list[dict[str, Any]] = []
        cards = soup.select("div[data-component-type='s-search-result']")
        for card in cards:
            asin = str(card.get("data-asin") or "").strip()
            title_node = card.select_one("h2 a span") or card.select_one("h2 span")
            title = title_node.get_text(" ", strip=True) if title_node else ""
            price_node = card.select_one(".a-price .a-offscreen") or card.select_one(".a-price")
            price = _num(price_node.get_text(" ", strip=True) if price_node else 0)
            if not title or price <= 0:
                continue
            link_node = card.select_one("h2 a[href]") or card.select_one("a.a-link-normal[href]")
            link = urljoin(page_url, str(link_node.get("href"))) if link_node else ""
            rating_node = card.select_one("span.a-icon-alt")
            rating = _parse_rating(rating_node.get_text(" ", strip=True) if rating_node else "")
            reviews_node = (
                card.select_one("a[href*='customerReviews'] span")
                or card.select_one("a[href*='product-reviews'] span")
                or card.select_one("span.a-size-base.s-underline-text")
            )
            records.append(
                {
                    "platform": self.key,
                    "product_id": asin or self._product_id(link, title),
                    "title": title,
                    "brand": title.split()[0] if title else "",
                    "url": link or page_url,
                    "market": self.cfg.get("market", "US"),
                    "currency": self.cfg.get("currency", "USD"),
                    "price_local": price,
                    "rating": rating,
                    "review_count": _int(reviews_node.get_text(" ", strip=True) if reviews_node else 0),
                    "seller": "",
                    "monthly_search_volume": 0,
                    "competitor_count": page_count or len(cards),
                    "ad_count": 1 if card.select_one("[data-component-type='sp-sponsored-result']") else 0,
                    "source_page": page_url,
                    "data_source": "public_live_page",
                    "collected_at": utils.utc_now_iso(),
                }
            )
        if records:
            return records
        return [self._record_from_jsonld(item, page_url) for item in _jsonld_products(html, page_url)]


class AliExpressAdapter(BasePlatformAdapter):
    def parse(self, html: str, page_url: str) -> list[dict[str, Any]]:
        soup = _soup(html)
        page_text = soup.get_text(" ", strip=True)
        count_match = re.search(r"([\d,]+)\s+(?:results|items)", page_text, re.IGNORECASE)
        page_count = _int(count_match.group(1)) if count_match else 0
        records: list[dict[str, Any]] = []
        cards = soup.select(
            "div.search-item-card-wrapper-gallery, a.search-item-card, div.product-card"
        )
        for card in cards:
            link_node = card.select_one("a[href*='/item/']") or card.select_one("a[href]")
            link = urljoin(page_url, str(link_node.get("href"))) if link_node else ""
            title_node = (
                card.select_one("h3")
                or card.select_one("[class*='title']")
                or card.select_one("a[title]")
            )
            title = title_node.get_text(" ", strip=True) if title_node else ""
            price_node = (
                card.select_one("[class*='price']")
                or card.select_one(".seo-optimized-price")
                or card.select_one(".product-price")
            )
            price = _num(price_node.get_text(" ", strip=True) if price_node else 0)
            if not title or price <= 0:
                continue
            rating_node = card.select_one("[class*='star']") or card.select_one("[class*='rating']")
            review_node = card.select_one("[class*='trade']") or card.select_one("[class*='review']")
            records.append(
                {
                    "platform": self.key,
                    "product_id": self._product_id(link, title),
                    "title": title,
                    "brand": "",
                    "url": link or page_url,
                    "market": self.cfg.get("market", "US"),
                    "currency": self.cfg.get("currency", "USD"),
                    "price_local": price,
                    "rating": _parse_rating(rating_node.get_text(" ", strip=True) if rating_node else ""),
                    "review_count": _int(review_node.get_text(" ", strip=True) if review_node else 0),
                    "seller": "",
                    "monthly_search_volume": 0,
                    "competitor_count": page_count or len(cards),
                    "ad_count": 0,
                    "source_page": page_url,
                    "data_source": "public_live_page",
                    "collected_at": utils.utc_now_iso(),
                }
            )
        if records:
            return records
        return [self._record_from_jsonld(item, page_url) for item in _jsonld_products(html, page_url)]


class Alibaba1688Adapter(BasePlatformAdapter):
    def parse(self, html: str, page_url: str) -> list[dict[str, Any]]:
        soup = _soup(html)
        cards = soup.select("div[data-offer-id], div.offer-list-row, div.sm-offer-item")
        records: list[dict[str, Any]] = []
        for card in cards:
            offer_id = str(card.get("data-offer-id") or "").strip()
            link_node = card.select_one("a[href*='detail.1688.com/offer']") or card.select_one("a[href]")
            link = urljoin(page_url, str(link_node.get("href"))) if link_node else ""
            title_node = (
                card.select_one(".title-text")
                or card.select_one("[class*='title']")
                or card.select_one("a[title]")
            )
            title = title_node.get_text(" ", strip=True) if title_node else ""
            if not title and link_node:
                title = str(link_node.get("title") or "")
            price_node = (
                card.select_one(".price")
                or card.select_one(".offer-price")
                or card.select_one("[class*='price']")
            )
            price = _num(price_node.get_text(" ", strip=True) if price_node else 0)
            if not title or price <= 0:
                continue
            seller_node = card.select_one(".company-name") or card.select_one(".shop-name")
            trade_node = card.select_one("[class*='trade']")
            moq_match = re.search(r"(\d+)\s*(?:件|个|台)", card.get_text(" ", strip=True))
            records.append(
                {
                    "platform": self.key,
                    "product_id": offer_id or self._product_id(link, title),
                    "title": title,
                    "brand": "",
                    "url": link or page_url,
                    "market": self.cfg.get("market", "US"),
                    "origin_country": "CN",
                    "currency": self.cfg.get("currency", "CNY"),
                    "price_local": price,
                    "rating": 0,
                    "review_count": _int(trade_node.get_text(" ", strip=True) if trade_node else 0),
                    "seller": seller_node.get_text(" ", strip=True) if seller_node else "",
                    "monthly_search_volume": 0,
                    "competitor_count": len(cards),
                    "ad_count": 0,
                    "moq": _int(moq_match.group(1)) if moq_match else 1,
                    "source_page": page_url,
                    "data_source": "public_live_page",
                    "collected_at": utils.utc_now_iso(),
                }
            )
        if records:
            return records
        return [self._record_from_jsonld(item, page_url) for item in _jsonld_products(html, page_url)]


ADAPTERS: dict[str, type[BasePlatformAdapter]] = {
    "amazon": AmazonAdapter,
    "aliexpress": AliExpressAdapter,
    "alibaba_1688": Alibaba1688Adapter,
}


def collect_live(
    config: ResearchConfig,
    *,
    limit: int,
    query: str | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    fetcher = PublicFetcher(config)
    records: list[dict[str, Any]] = []
    failures: list[dict[str, str]] = []
    try:
        for key, platform_cfg in config.platform_items():
            adapter_cls = ADAPTERS.get(key)
            if adapter_cls is None:
                failures.append({"platform": key, "error": "未实现公开页面适配器"})
                continue
            try:
                items = adapter_cls(key, platform_cfg, fetcher).crawl(limit=limit, query=query)
                records.extend(items)
                logger.info("live 抓取成功: %s，%d 条", platform_cfg.get("label", key), len(items))
            except (CrawlError, ParseError) as exc:
                logger.warning("live 抓取失败: %s：%s", platform_cfg.get("label", key), exc)
                failures.append({"platform": key, "error": str(exc)})
    finally:
        fetcher.close()
    return records, failures


def load_fixture(config: ResearchConfig) -> list[dict[str, Any]]:
    fixture_path = config.resolve("fixture.path")
    payload = utils.read_json(fixture_path)
    if not isinstance(payload, dict) or not isinstance(payload.get("records"), list):
        raise ResearchError(f"公开快照格式错误: {fixture_path}")
    captured_at = str(payload.get("captured_at") or utils.utc_now_iso())
    records: list[dict[str, Any]] = []
    for item in payload["records"]:
        if not isinstance(item, dict):
            continue
        record = dict(item)
        record.setdefault("data_source", "offline_public_snapshot")
        record.setdefault("collected_at", captured_at)
        records.append(record)
    logger.info("已加载离线公开快照: %s，%d 条", fixture_path, len(records))
    return records


def collect_public_data(
    config: ResearchConfig,
    *,
    mode: str,
    limit: int,
    query: str | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    mode = mode.lower()
    if mode == "offline":
        return load_fixture(config), {"mode": "offline", "live_failures": []}
    if mode == "live":
        records, failures = collect_live(config, limit=limit, query=query)
        if not records:
            raise CrawlError(
                "live 模式未获取到任何公开商品；请检查网络、robots.txt 或页面结构，"
                "不要通过验证码/Cookie/代理绕过访问控制。"
            )
        return records, {"mode": "live", "live_failures": failures}
    if mode == "auto":
        try:
            records, failures = collect_live(config, limit=limit, query=query)
            if records:
                return records, {"mode": "auto_live", "live_failures": failures}
        except CrawlError as exc:
            logger.warning("auto 模式 live 失败，回退快照：%s", exc)
        return load_fixture(config), {"mode": "auto_fixture", "live_failures": []}
    raise ResearchError(f"未知 mode: {mode}")


# ---------------------------------------------------------------------------
# Excel 输出
# ---------------------------------------------------------------------------
# (字段名, 中文表头, 列宽, 数字格式)
CORE_COLUMNS: list[tuple[str, str, int, str]] = [
    ("rank", "排名", 8, "0"),
    ("platform_label", "平台", 11, "@"),
    ("title", "商品标题", 48, "@"),
    ("brand", "品牌/供应商", 16, "@"),
    ("sku", "SKU", 18, "@"),
    ("url", "商品链接", 30, "@"),
    ("market", "目标市场", 10, "@"),
    ("capacity_wh", "容量(Wh)", 10, "#,##0"),
    ("output_w", "输出功率(W)", 11, "#,##0"),
    ("weight_kg", "重量(kg)", 10, "0.00"),
    ("battery_type", "电芯类型", 12, "@"),
    ("solar_input_w", "太阳能输入(W)", 13, "#,##0"),
    ("rating", "评分", 8, "0.0"),
    ("review_count", "评论/成交数", 12, "#,##0"),
    ("seller", "卖家/供应商", 20, "@"),
    ("price_local", "公开售价(原币)", 15, "#,##0.00"),
    ("currency", "币种", 8, "@"),
    ("price_usd", "公开价(USD)", 13, "$#,##0.00"),
    ("retail_price_usd", "零售基准价(USD)", 15, "$#,##0.00"),
    ("supply_cost_usd", "货源/估算成本(USD)", 17, "$#,##0.00"),
    ("landed_cost_usd", "到岸成本(USD)", 14, "$#,##0.00"),
    ("platform_fee_usd", "平台佣金(USD)", 13, "$#,##0.00"),
    ("fulfillment_cost_usd", "履约成本(USD)", 13, "$#,##0.00"),
    ("gross_profit_usd", "单件毛利(USD)", 14, "$#,##0.00"),
    ("gross_margin_pct", "毛利率(%)", 11, "0.0"),
    ("ad_cost_est_usd", "广告估算(USD)", 13, "$#,##0.00"),
    ("contribution_margin_pct", "贡献毛利率(%)", 14, "0.0"),
    ("monthly_search_volume", "月搜索量", 11, "#,##0"),
    ("monthly_units", "预估月销量", 11, "#,##0"),
    ("market_capacity_usd", "市场容量(USD/月)", 16, "$#,##0"),
    ("addressable_revenue_usd", "可获取容量(USD/月)", 17, "$#,##0"),
    ("competitor_count", "竞品数量", 10, "#,##0"),
    ("ad_count", "广告位数", 10, "#,##0"),
    ("competition_score", "竞争度(越低越好)", 15, "0.0"),
    ("margin_score", "毛利得分", 10, "0.0"),
    ("market_score", "市场得分", 10, "0.0"),
    ("competition_opportunity_score", "竞争机会分", 12, "0.0"),
    ("lightweight_score", "轻量得分", 10, "0.0"),
    ("data_quality_score", "数据质量分", 11, "0.0"),
    ("opportunity_score", "综合机会分", 12, "0.0"),
    ("recommendation", "推荐等级", 14, "@"),
    ("gate_status", "门槛状态", 30, "@"),
    ("compliance_note", "合规提示", 46, "@"),
    ("target_channel", "目标渠道", 20, "@"),
    ("cost_basis", "成本口径", 36, "@"),
    ("suggested_action", "建议动作", 42, "@"),
    ("moq", "MOQ", 8, "#,##0"),
    ("price_per_wh_usd", "USD/Wh", 10, "0.000"),
    ("source_page", "公开来源页", 36, "@"),
    ("captured_at", "采集时间", 20, "@"),
    ("record_id", "原始记录ID", 18, "@"),
    ("data_source", "数据来源", 22, "@"),
    ("demand_source", "需求口径", 16, "@"),
    ("weight_source", "重量来源", 12, "@"),
    ("retail_benchmark_source", "零售基准来源", 18, "@"),
    ("price_band", "价格带", 10, "@"),
]
COLUMN_KEYS = [item[0] for item in CORE_COLUMNS]
MONEY_KEYS = {
    "price_usd",
    "retail_price_usd",
    "supply_cost_usd",
    "landed_cost_usd",
    "platform_fee_usd",
    "fulfillment_cost_usd",
    "gross_profit_usd",
    "ad_cost_est_usd",
    "market_capacity_usd",
    "addressable_revenue_usd",
}


class ExcelExporter:
    """生成带格式、筛选、指标说明和数据质量页的 Excel。"""

    def __init__(self, config: ResearchConfig) -> None:
        self.config = config

    def export(
        self,
        *,
        high_rows: list[dict[str, Any]],
        all_rows: list[dict[str, Any]],
        summary: dict[str, Any],
        output_path: Path,
    ) -> Path:
        if openpyxl is None:
            raise ResearchError("生成 Excel 需要 openpyxl：pip install openpyxl")
        from openpyxl import Workbook
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "高潜力选品"
        self._write_data_sheet(
            ws,
            high_rows,
            metadata=summary,
            title="户外露营轻量化储能 · 高潜力选品表（100-500Wh）",
            with_meta=True,
        )
        self._apply_highlights(ws, len(high_rows), ColorScaleRule)

        ws_all = wb.create_sheet("全部清洗数据")
        self._write_data_sheet(
            ws_all,
            all_rows,
            metadata=summary,
            title="全部清洗数据",
            with_meta=False,
        )
        self._apply_highlights(ws_all, len(all_rows), ColorScaleRule, header_row=1)

        ws_metric = wb.create_sheet("指标说明")
        self._write_metric_sheet(ws_metric, summary, Alignment, Border, Font, PatternFill, Side)

        ws_params = wb.create_sheet("筛选参数")
        self._write_params_sheet(ws_params, Alignment, Border, Font, PatternFill, Side)

        ws_quality = wb.create_sheet("数据质量")
        self._write_quality_sheet(ws_quality, summary, Alignment, Border, Font, PatternFill, Side)

        wb.properties.title = "户外露营轻量化储能高潜力选品表"
        wb.properties.creator = "Outdoor ESS AI product_research.py"
        wb.properties.description = "公开数据选品分析；不接入平台店铺 API。"
        utils.ensure_dir(output_path.parent)
        wb.save(output_path)
        logger.info("Excel 已生成: %s", output_path)
        return output_path

    def _write_data_sheet(
        self,
        ws: Any,
        rows: list[dict[str, Any]],
        *,
        metadata: dict[str, Any],
        title: str,
        with_meta: bool,
    ) -> None:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        header_row = 6 if with_meta else 1
        if with_meta:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
            ws["A1"] = title
            ws["A1"].font = Font(bold=True, size=16, color="1F4E78")
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=16)
            ws["A2"] = (
                f"运行：{metadata.get('run_id', '-')}  |  数据模式：{metadata.get('mode', '-')}  |  "
                f"生成：{metadata.get('generated_at', '-')}  |  样本：{metadata.get('cleaned_records', 0)} 条 / "
                f"高潜力：{metadata.get('high_potential_records', 0)} 条"
            )
            ws["A2"].font = Font(size=9, color="595959")
            ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=18)
            ws["A3"] = (
                f"筛选：容量 {metadata.get('filters', {}).get('capacity_wh', '-')} Wh ｜ "
                f"重量≤{metadata.get('filters', {}).get('max_weight_kg', '-')}kg ｜ "
                f"毛利率≥{metadata.get('filters', {}).get('min_gross_margin_pct', '-')}% ｜ "
                f"市场容量≥${metadata.get('filters', {}).get('min_market_capacity_usd', 0):,.0f}/月 ｜ "
                f"竞争度≤{metadata.get('filters', {}).get('max_competition_score', '-')}"
            )
            ws["A3"].font = Font(size=9, color="595959")
            ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=20)
            ws["A4"] = (
                "口径：毛利=零售价-到岸成本-平台佣金-履约成本；市场容量=月搜索量×转化率×零售价；"
                "竞争度越高越激烈。公开数据/模拟数据不构成认证或投资建议。"
            )
            ws["A4"].font = Font(size=9, color="7F6000")
            ws["A4"].fill = PatternFill("solid", fgColor="FFF2CC")
            for row in range(1, 5):
                ws.row_dimensions[row].height = 24 if row == 1 else 20

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True, size=10)
        border = Border(
            left=Side(style="thin", color="D9E2F3"),
            right=Side(style="thin", color="D9E2F3"),
            top=Side(style="thin", color="D9E2F3"),
            bottom=Side(style="thin", color="D9E2F3"),
        )
        for col_index, (_key, header, width, _fmt) in enumerate(CORE_COLUMNS, start=1):
            cell = ws.cell(row=header_row, column=col_index, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            ws.column_dimensions[get_column_letter(col_index)].width = width
        ws.row_dimensions[header_row].height = 38

        if not rows:
            ws.merge_cells(
                start_row=header_row + 1,
                start_column=1,
                end_row=header_row + 1,
                end_column=8,
            )
            ws.cell(row=header_row + 1, column=1, value="当前参数下无符合全部门槛的 SKU，请放宽阈值后重跑。")
            ws.cell(row=header_row + 1, column=1).font = Font(color="C00000", italic=True)
            return

        for row_index, row in enumerate(rows, start=header_row + 1):
            for col_index, (key, _header, _width, fmt) in enumerate(CORE_COLUMNS, start=1):
                value = row.get(key)
                if isinstance(value, (list, dict)):
                    value = json.dumps(value, ensure_ascii=False)
                cell = ws.cell(row=row_index, column=col_index, value=value)
                cell.border = border
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=key in {"title", "compliance_note", "suggested_action", "gate_status", "cost_basis"},
                )
                if key in MONEY_KEYS or key == "price_local":
                    cell.number_format = "$#,##0.00" if key in MONEY_KEYS else "#,##0.00"
                elif fmt:
                    cell.number_format = fmt
                if key in {"url", "source_page"} and value:
                    cell.hyperlink = str(value)
                    cell.font = Font(color="0563C1", underline="single")
            ws.row_dimensions[row_index].height = 34

        last_col = get_column_letter(len(CORE_COLUMNS))
        ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + len(rows)}"
        ws.freeze_panes = f"A{header_row + 1}"
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    @staticmethod
    def _apply_highlights(ws: Any, row_count: int, color_scale_cls: Any, *, header_row: int = 6) -> None:
        if row_count <= 0:
            return
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.utils import get_column_letter

        positions = {key: index for index, (key, *_rest) in enumerate(CORE_COLUMNS, start=1)}
        last_row = header_row + row_count
        for key, reverse in (("opportunity_score", False), ("gross_margin_pct", False), ("competition_score", True)):
            col = get_column_letter(positions[key])
            if reverse:
                rule = ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="63BE7B",
                    mid_type="num",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="num",
                    end_value=100,
                    end_color="F8696B",
                )
            else:
                rule = ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="F8696B",
                    mid_type="num",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="num",
                    end_value=100,
                    end_color="63BE7B",
                )
            ws.conditional_formatting.add(f"{col}{header_row + 1}:{col}{last_row}", rule)

    def _write_metric_sheet(self, ws: Any, summary: dict[str, Any], *_: Any) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        ws.title = "指标说明"
        ws.append(["指标", "字段", "计算公式", "业务解释", "当前门槛"])
        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in ws[1]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        metrics = [
            ("单件毛利", "gross_profit_usd", "零售基准价 - 到岸成本 - 平台佣金 - 履约成本", "单件商品的可变成本后利润，不含广告与固定费用。", "-"),
            ("毛利率", "gross_margin_pct", "单件毛利 ÷ 零售基准价 × 100%", "越高越好；低于门槛不进入高潜表。", f"≥{summary.get('filters', {}).get('min_gross_margin_pct', 25)}%"),
            ("到岸成本", "landed_cost_usd", "货源/估算成本 + 头程 + 关税 + 包装", "零售平台成本按公开价×成本率估算；1688 按公开批发价。", "-"),
            ("市场容量", "market_capacity_usd", "月搜索量 × 平台转化率 × 零售基准价", "目标市场月度可寻址销售额，不代表本 SKU 实际销量。", f"≥${summary.get('filters', {}).get('min_market_capacity_usd', 120000):,.0f}"),
            ("可获取容量", "addressable_revenue_usd", "市场容量 × 可获取份额", "按配置的 addressable_share 估算，默认 15%。", "-"),
            ("竞争度", "competition_score", "0.35×列表密度 + 0.30×评论壁垒 + 0.20×广告压力 + 0.15×价格离散", "0-100，越高越激烈；因此反向进入综合分。", f"≤{summary.get('filters', {}).get('max_competition_score', 70)}"),
            ("综合机会分", "opportunity_score", "0.35×毛利得分 + 0.25×市场得分 + 0.20×竞争机会 + 0.10×轻量得分 + 0.10×数据质量", "0-100，用于排序；仅作决策辅助。", f"≥{summary.get('filters', {}).get('min_opportunity_score', 65)}"),
            ("轻量化得分", "lightweight_score", "100 - max(0, 重量-2kg)×7.5", "越轻越高，最高 100。", "-"),
            ("数据质量分", "data_quality_score", "100 - 缺失字段惩罚 - 估算重量惩罚", "公开页面字段缺失会降低可信度。", "-"),
        ]
        for item in metrics:
            ws.append(list(item))
        widths = [16, 24, 52, 60, 22]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + index)].width = width
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"

    def _write_params_sheet(self, ws: Any, *_: Any) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        ws.title = "筛选参数"
        ws.append(["参数", "当前值", "说明"])
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
        for key, value in _flatten(self.config.data).items():
            rendered = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
            ws.append([key, rendered, "修改 config/product_research.yaml 或使用 CLI 参数覆盖"])
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 46
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"

    def _write_quality_sheet(self, ws: Any, summary: dict[str, Any], *_: Any) -> None:
        from openpyxl.styles import Font, PatternFill

        ws.title = "数据质量"
        ws.append(["类别", "指标", "数值"])
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
        quality = summary.get("quality", {})
        rows = [
            ("总量", "原始记录", quality.get("raw_records", 0)),
            ("总量", "清洗后记录", quality.get("clean_records", 0)),
            ("总量", "高潜力记录", summary.get("high_potential_records", 0)),
            ("总量", "平均毛利率(%)", quality.get("avg_margin_pct", 0)),
            ("总量", "平均竞争度", quality.get("avg_competition", 0)),
            ("总量", "市场中位容量(USD/月)", quality.get("median_market_capacity", 0)),
        ]
        for key, value in (quality.get("by_platform_raw") or {}).items():
            rows.append(("平台原始", key, value))
        for key, value in (quality.get("by_platform_clean") or {}).items():
            rows.append(("平台清洗后", key, value))
        for key, value in (quality.get("dropped_by_reason") or {}).items():
            rows.append(("剔除原因", key, value))
        for key, value in (quality.get("benchmark") or {}).items():
            rows.append(("零售基准价", f"{key} Wh", value))
        for failure in summary.get("live_failures", []):
            rows.append(("live 失败", failure.get("platform", ""), failure.get("error", "")))
        for item in rows:
            ws.append(list(item))
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["C"].width = 80
        ws.freeze_panes = "A2"


def _flatten(value: Any, prefix: str = "") -> dict[str, Any]:
    result: dict[str, Any] = {}
    if isinstance(value, dict):
        for key, child in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            result.update(_flatten(child, path))
    else:
        result[prefix] = value
    return result


def _display_path(root: Path, path: Path) -> str:
    """优先返回相对项目根目录的路径；外部输出目录返回绝对路径。"""

    try:
        return str(path.resolve().relative_to(root.resolve()))
    except ValueError:
        return str(path.resolve())


# ---------------------------------------------------------------------------
# 引擎与 CLI
# ---------------------------------------------------------------------------
def apply_overrides(
    config: ResearchConfig,
    *,
    mode: str | None = None,
    output_dir: str | None = None,
    fixture: str | None = None,
    top_n: int | None = None,
    min_wh: float | None = None,
    max_wh: float | None = None,
    max_weight: float | None = None,
    min_margin: float | None = None,
    min_market_capacity: float | None = None,
    max_competition: float | None = None,
    min_score: float | None = None,
    min_price: float | None = None,
    max_price: float | None = None,
) -> ResearchConfig:
    data = copy.deepcopy(config.data)
    if mode:
        data.setdefault("crawler", {})["mode"] = mode
    if fixture:
        data.setdefault("fixture", {})["path"] = fixture
    if output_dir:
        data.setdefault("output", {})["dir"] = output_dir
    filters = data.setdefault("filters", {})
    if min_wh is not None or max_wh is not None:
        current = list(filters.get("capacity_wh", [100, 500]))
        filters["capacity_wh"] = [
            float(min_wh if min_wh is not None else current[0]),
            float(max_wh if max_wh is not None else current[1]),
        ]
    if max_weight is not None:
        filters["max_weight_kg"] = float(max_weight)
    if min_price is not None or max_price is not None:
        current = list(filters.get("price_usd", [49, 1500]))
        filters["price_usd"] = [
            float(min_price if min_price is not None else current[0]),
            float(max_price if max_price is not None else current[1]),
        ]
    thresholds = data.setdefault("opportunity", {}).setdefault("thresholds", {})
    if min_margin is not None:
        thresholds["min_gross_margin_pct"] = float(min_margin)
    if min_market_capacity is not None:
        thresholds["min_market_capacity_usd"] = float(min_market_capacity)
    if max_competition is not None:
        thresholds["max_competition_score"] = float(max_competition)
    if min_score is not None:
        thresholds["min_opportunity_score"] = float(min_score)
    if top_n is not None:
        data.setdefault("opportunity", {})["top_n"] = int(top_n)
    return ResearchConfig(root=config.root, path=config.path, data=data)


class ProductResearchEngine:
    """采集 → 清洗 → 指标 → Excel/CSV/JSON 的独立编排器。"""

    def __init__(self, config: ResearchConfig) -> None:
        self.config = config
        self.output_dir = config.resolve("output.dir")
        self.cleaner = ProductCleaner(config)
        self.metrics = SelectionMetrics(config)
        self.exporter = ExcelExporter(config)

    def run(
        self,
        *,
        mode: str | None = None,
        limit: int | None = None,
        query: str | None = None,
    ) -> dict[str, Any]:
        utils.ensure_dir(self.output_dir)
        selected_mode = str(mode or self.config.get("crawler.mode", "offline")).lower()
        limit = int(limit or self.config.get("crawler.max_results_per_platform", 30))
        run_id = f"research_{utils.timestamp()}"
        generated_at = utils.utc_now_iso()

        raw_records, meta = collect_public_data(
            self.config, mode=selected_mode, limit=limit, query=query
        )
        clean_rows, quality = self.cleaner.clean(raw_records)
        all_rows, high_rows, metric_stats = self.metrics.enrich(clean_rows)

        thresholds = dict(self.config.get("opportunity.thresholds", {}) or {})
        filters = {
            "capacity_wh": [float(x) for x in self.config.get("filters.capacity_wh", [100, 500])],
            "max_weight_kg": float(self.config.get("filters.max_weight_kg", 15)),
            "price_usd": [float(x) for x in self.config.get("filters.price_usd", [49, 1500])],
            "min_gross_margin_pct": float(thresholds.get("min_gross_margin_pct", 25)),
            "min_market_capacity_usd": float(thresholds.get("min_market_capacity_usd", 120000)),
            "max_competition_score": float(thresholds.get("max_competition_score", 70)),
            "max_weight_threshold_kg": float(thresholds.get("max_weight_kg", 12)),
            "min_opportunity_score": float(thresholds.get("min_opportunity_score", 65)),
            "top_n": int(self.config.get("opportunity.top_n", 30)),
        }

        raw_path = utils.write_json(
            self.output_dir / str(self.config.get("output.raw_name", "raw_public_snapshot.json")),
            {
                "run_id": run_id,
                "generated_at": generated_at,
                "mode": meta.get("mode", selected_mode),
                "source_policy": "public-visible-only; no login; no private API; no captcha bypass",
                "live_failures": meta.get("live_failures", []),
                "records": raw_records,
            },
        )
        csv_path = utils.write_csv(
            self.output_dir / str(self.config.get("output.csv_name", "cleaned_products.csv")),
            all_rows,
            fieldnames=COLUMN_KEYS,
        )
        excel_path = self.output_dir / str(
            self.config.get("output.excel_name", "high_potential_products.xlsx")
        )

        summary: dict[str, Any] = {
            "run_id": run_id,
            "generated_at": generated_at,
            "mode": meta.get("mode", selected_mode),
            "data_policy": "public-visible-only; no login; no private API; no captcha bypass",
            "notice": "公开数据/模拟数据选品结果，不构成认证、法律或投资建议。",
            "raw_records": quality.get("raw_records", len(raw_records)),
            "cleaned_records": len(all_rows),
            "high_potential_records": len(high_rows),
            "filters": filters,
            "quality": {**quality, **metric_stats},
            "live_failures": meta.get("live_failures", []),
            "top5": [
                {
                    "rank": row["rank"],
                    "platform": row["platform_label"],
                    "title": row["title"],
                    "capacity_wh": row["capacity_wh"],
                    "gross_margin_pct": row["gross_margin_pct"],
                    "market_capacity_usd": row["market_capacity_usd"],
                    "competition_score": row["competition_score"],
                    "opportunity_score": row["opportunity_score"],
                }
                for row in (high_rows or all_rows)[:5]
            ],
            "outputs": {
                "excel": _display_path(self.config.root, excel_path),
                "csv": _display_path(self.config.root, csv_path),
                "raw": _display_path(self.config.root, raw_path),
            },
        }
        self.exporter.export(
            high_rows=high_rows,
            all_rows=all_rows,
            summary=summary,
            output_path=excel_path,
        )
        summary_path = self.output_dir / str(
            self.config.get("output.summary_name", "research_summary.json")
        )
        summary["outputs"]["summary"] = _display_path(self.config.root, summary_path)
        utils.write_json(summary_path, summary)
        logger.info(
            "选品完成: 原始 %d / 清洗 %d / 高潜力 %d",
            summary["raw_records"],
            summary["cleaned_records"],
            summary["high_potential_records"],
        )
        return summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="product_research",
        description="公开数据选品：100-500Wh 户外露营轻量化储能",
    )
    parser.add_argument("--config", default=str(DEFAULT_CONFIG), help="选品配置文件")
    parser.add_argument("--mode", choices=["offline", "live", "auto"], help="数据模式")
    parser.add_argument("--fixture", help="离线公开快照 JSON 路径")
    parser.add_argument("--output-dir", help="输出目录，默认 outputs/research")
    parser.add_argument("--limit", type=int, help="每个平台最多抓取条数（live）")
    parser.add_argument("--query", help="live 搜索关键词，默认使用平台配置")
    parser.add_argument("--top-n", type=int, help="高潜表最多输出条数")
    parser.add_argument("--min-wh", type=float, help="最小容量 Wh")
    parser.add_argument("--max-wh", type=float, help="最大容量 Wh")
    parser.add_argument("--max-weight", type=float, help="最大重量 kg")
    parser.add_argument("--min-price", type=float, help="最低公开价 USD")
    parser.add_argument("--max-price", type=float, help="最高公开价 USD")
    parser.add_argument("--min-margin", type=float, help="最低毛利率 %%")
    parser.add_argument("--min-market-capacity", type=float, help="最低市场容量 USD/月")
    parser.add_argument("--max-competition", type=float, help="最高竞争度 0-100")
    parser.add_argument("--min-score", type=float, help="最低综合机会分")
    parser.add_argument("--log-level", default="INFO", help="日志级别")
    parser.add_argument("--strict", action="store_true", help="无高潜力 SKU 时返回非零退出码")
    return parser


def _configure_stdio() -> None:
    """Windows GBK 控制台下强制 UTF-8，避免中文和符号输出失败。"""

    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if callable(reconfigure):
            reconfigure(encoding="utf-8", errors="replace")


def main(argv: Iterable[str] | None = None) -> int:
    _configure_stdio()
    args = build_parser().parse_args(list(argv) if argv is not None else None)
    try:
        config = load_research_config(args.config)
        config = apply_overrides(
            config,
            mode=args.mode,
            output_dir=args.output_dir,
            fixture=args.fixture,
            top_n=args.top_n,
            min_wh=args.min_wh,
            max_wh=args.max_wh,
            max_weight=args.max_weight,
            min_margin=args.min_margin,
            min_market_capacity=args.min_market_capacity,
            max_competition=args.max_competition,
            min_score=args.min_score,
            min_price=args.min_price,
            max_price=args.max_price,
        )
    except (ResearchError, utils.UtilError) as exc:
        print(f"✗ 配置加载失败：{exc}", file=sys.stderr)
        return 1
    output_dir = config.resolve("output.dir")
    utils.ensure_dir(output_dir)
    utils.setup_logging(
        level=args.log_level,
        log_file=output_dir / "product_research.log",
        console=True,
        force=True,
    )
    try:
        summary = ProductResearchEngine(config).run(
            mode=args.mode,
            limit=args.limit,
            query=args.query,
        )
    except (ResearchError, CrawlError, utils.UtilError) as exc:
        logger.error("选品失败: %s", exc)
        print(f"✗ 选品失败：{exc}", file=sys.stderr)
        return 1

    print("")
    print("=" * 72)
    print(f"✓ 选品完成：{summary['run_id']}")
    print(f"  模式      : {summary['mode']}")
    print(f"  原始/清洗 : {summary['raw_records']} / {summary['cleaned_records']}")
    print(f"  高潜力 SKU: {summary['high_potential_records']}")
    print(
        f"  平均毛利  : {summary['quality'].get('avg_margin_pct', 0):.1f}%  |  "
        f"平均竞争度: {summary['quality'].get('avg_competition', 0):.1f}"
    )
    print(f"  Excel     : {summary['outputs']['excel']}")
    print(f"  CSV       : {summary['outputs']['csv']}")
    print(f"  JSON      : {summary['outputs']['summary']}")
    print("=" * 72)
    if summary["high_potential_records"] == 0 and args.strict:
        return 2
    return 0


if __name__ == "__main__":
    from error_guard import run_cli

    raise SystemExit(run_cli(main, "product_research"))
