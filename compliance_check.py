#!/usr/bin/env python3
"""欧美储能产品合规风险筛查。

输入：
    - 第二步选品表（SKU / 商品名）；
    - 模拟产品合规主数据（电压、频率、电池、无线、认证资料等）。

输出：
    - 结构化 Excel：汇总、不符合项、产品参数、标准依据；
    - JSON 摘要。

说明：
    CE 是适用指令/法规下的符合性评估与自我声明，不是一张“CE 证书”；
    FCC 按设备类型区分 SDoC 与认证（FCC ID）。本脚本只做规则初筛，
    不替代测试实验室、认证机构、律师或海关/平台审核。
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parent
VENDOR_DIR = PROJECT_ROOT / ".vendor"
if VENDOR_DIR.is_dir() and str(VENDOR_DIR) not in sys.path:
    sys.path.insert(0, str(VENDOR_DIR))

import utils  # noqa: E402

logger = utils.get_logger("compliance_check")
DEFAULT_SELECTION = PROJECT_ROOT / "outputs" / "research" / "high_potential_products.xlsx"
DEFAULT_MASTER = PROJECT_ROOT / "data" / "mock" / "product_compliance_master.json"
DEFAULT_RULES = PROJECT_ROOT / "config" / "compliance_rules.yaml"
DEFAULT_OUTPUT = PROJECT_ROOT / "outputs" / "compliance" / "compliance_check.xlsx"
DEFAULT_JSON = PROJECT_ROOT / "outputs" / "compliance" / "compliance_summary.json"
SEVERITY_ORDER = {"high": 3, "medium": 2, "low": 1}
STATUS_LABEL = {"high": "不符合", "medium": "待补充", "low": "建议"}


class ComplianceError(RuntimeError):
    """合规筛查异常。"""


@dataclass(frozen=True)
class Rules:
    root: Path
    data: dict[str, Any]

    @property
    def weights(self) -> dict[str, float]:
        return {str(k): float(v) for k, v in (self.data.get("severity_weights") or {}).items()}

    @property
    def standards(self) -> dict[str, dict[str, Any]]:
        return {
            str(item["id"]): dict(item)
            for item in (self.data.get("standards") or [])
            if isinstance(item, dict) and item.get("id")
        }

    def get(self, dotted: str, default: Any = None) -> Any:
        return utils.deep_get(self.data, dotted, default)


def load_rules(path: str | Path = DEFAULT_RULES) -> Rules:
    target = Path(path).expanduser().resolve()
    data = utils.read_yaml(target)
    if not isinstance(data, dict) or not data.get("standards"):
        raise ComplianceError(f"合规规则格式错误: {target}")
    return Rules(root=PROJECT_ROOT, data=data)


def _num(value: Any, default: float = 0.0) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value or ""))
    return float(match.group(0)) if match else default


def _load_selection(path: str | Path) -> dict[str, dict[str, Any]]:
    source = Path(path).expanduser().resolve()
    if not source.exists():
        logger.warning("选品表不存在，仅使用合规主数据: %s", source)
        return {}
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ComplianceError("读取选品表需要 openpyxl") from exc
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb["高潜力选品"] if "高潜力选品" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers: list[str] = []
    header_row = 1
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        values = [str(cell or "").strip() for cell in row]
        if "SKU" in values or "sku" in values:
            headers = values
            header_row = row_index
            break
    result: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        record = dict(zip(headers, row, strict=False))
        sku = str(record.get("SKU") or record.get("sku") or "").strip()
        if not sku:
            continue
        result[sku] = {
            "name": str(record.get("商品标题") or record.get("title") or ""),
            "platform": str(record.get("平台") or record.get("platform_label") or ""),
            "market": str(record.get("目标市场") or record.get("market") or ""),
            "price_usd": _num(record.get("零售基准价(USD)") or record.get("retail_price_usd")),
        }
    wb.close()
    return result


def load_master(path: str | Path = DEFAULT_MASTER, selection: dict[str, dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    master_path = Path(path).expanduser().resolve()
    payload = utils.read_json(master_path)
    if not isinstance(payload, dict) or not isinstance(payload.get("products"), list):
        raise ComplianceError(f"产品合规主数据格式错误: {master_path}")
    selection = selection or {}
    products: list[dict[str, Any]] = []
    for item in payload["products"]:
        if not isinstance(item, dict) or not item.get("sku"):
            continue
        product = dict(item)
        extra = selection.get(str(product["sku"]), {})
        product["name"] = product.get("name") or extra.get("name") or product.get("model") or product["sku"]
        product["platform_label"] = extra.get("platform", "")
        product["output_w"] = product.get("output_w") or product.get("rated_power_w")
        product["target_markets"] = [
            str(m).upper() for m in (product.get("target_markets") or ["US", "EU"])
        ]
        products.append(product)
    if not products:
        raise ComplianceError("合规主数据为空")
    return products


@dataclass
class Finding:
    sku: str
    product_name: str
    market: str
    standard_id: str
    standard_label: str
    requirement: str
    observed: str
    severity: str
    suggestion: str
    evidence: str
    status: str
    finding_id: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "finding_id": self.finding_id,
            "sku": self.sku,
            "product": self.product_name,
            "market": self.market,
            "standard_id": self.standard_id,
            "standard": self.standard_label,
            "requirement": self.requirement,
            "observed": self.observed,
            "severity": self.severity,
            "status": self.status,
            "suggestion": self.suggestion,
            "evidence_needed": self.evidence,
        }


class ComplianceChecker:
    def __init__(self, rules: Rules) -> None:
        self.rules = rules
        self.weights = rules.weights or {"high": 25, "medium": 10, "low": 5}
        self.standards = rules.standards
        self.thresholds = rules.data.get("thresholds", {}) or {}
        self.manual_languages = rules.data.get("manual_languages", {}) or {}

    def check(self, products: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[Finding]]:
        rows: list[dict[str, Any]] = []
        findings: list[Finding] = []
        for product in products:
            product_findings: list[Finding] = []
            markets = [str(m).upper() for m in product.get("target_markets", [])]
            for market in markets:
                if market == "US":
                    product_findings.extend(self._check_us(product))
                elif market in {"EU", "DE", "FR"}:
                    product_findings.extend(self._check_eu(product))
            product_findings.extend(self._check_general(product))
            findings.extend(product_findings)
            rows.append(self._summary_row(product, product_findings, markets))
        findings.sort(key=lambda item: (item.sku, -SEVERITY_ORDER.get(item.severity, 0), item.market))
        for index, finding in enumerate(findings, start=1):
            finding.finding_id = f"F-{index:03d}"
        return rows, findings

    # ------------------------------------------------------------------
    def _add(
        self,
        product: dict[str, Any],
        findings: list[Finding],
        *,
        market: str,
        standard_id: str,
        requirement: str,
        observed: str,
        severity: str,
        suggestion: str,
    ) -> None:
        standard = self.standards.get(standard_id, {})
        findings.append(
            Finding(
                sku=str(product["sku"]),
                product_name=str(product.get("name") or product.get("model") or ""),
                market=market,
                standard_id=standard_id,
                standard_label=str(standard.get("label") or standard_id),
                requirement=requirement,
                observed=observed,
                severity=severity,
                status=STATUS_LABEL.get(severity, "待确认"),
                suggestion=suggestion,
                evidence=str(standard.get("evidence") or "测试报告/技术文件/符合性声明"),
            )
        )

    def _check_eu(self, product: dict[str, Any]) -> list[Finding]:
        findings: list[Finding] = []
        voltage = _num(product.get("output_voltage_v"))
        frequency = _num(product.get("output_frequency_hz"))
        low, high = [float(x) for x in self.thresholds.get("eu_voltage_v", [220, 240])]
        eu_freq = float(self.thresholds.get("eu_frequency_hz", 50))
        if not low <= voltage <= high:
            self._add(
                product,
                findings,
                market="EU",
                standard_id="EU_LVD",
                requirement="欧盟市场输出额定电压应为 220-240V",
                observed=f"当前输出 {voltage:g}V / {frequency:g}Hz",
                severity="high",
                suggestion="为欧盟开发 230V/50Hz 版本，或明确不在 EU 销售该 SKU；同步更新铭牌、包装和 Listing。",
            )
        if frequency != eu_freq:
            self._add(
                product,
                findings,
                market="EU",
                standard_id="EU_LVD",
                requirement="欧盟市场输出频率应为 50Hz",
                observed=f"当前频率 {frequency:g}Hz",
                severity="high",
                suggestion="更换 EU 版本逆变器/主板并重新做 LVD/EMC 测试。",
            )
        self._require_eu_doc(
            product,
            findings,
            "EU_LVD",
            "has_lvd_report",
            "LVD 测试报告/技术文件",
            "high",
            "补齐 LVD 测试报告、风险评估、技术文件与 DoC。",
        )
        self._require_eu_doc(
            product,
            findings,
            "EU_EMC",
            "has_emc_report",
            "EMC 测试报告",
            "high",
            "补齐 EMC 发射/抗扰测试，并把测试报告纳入技术文件。",
        )
        self._require_eu_doc(
            product,
            findings,
            "EU_ROHS",
            "has_rohs_report",
            "RoHS 均质材料检测/声明",
            "high",
            "补齐 RoHS 10 项检测或供应链声明，保留物料清单与报告。",
        )
        self._require_eu_doc(
            product,
            findings,
            "EU_WEEE",
            "has_weee_registration",
            "WEEE 注册与回收方案",
            "medium",
            "完成目标成员国 WEEE 注册，增加带叉垃圾桶标识。",
        )
        self._require_eu_doc(
            product,
            findings,
            "EU_BATTERY",
            "has_battery_regulation_doc",
            "EU Battery Regulation 文件",
            "high",
            "建立电池碳足迹/回收料/标签/护照路线图，并保留电池安全与供应链文件。",
        )
        self._require_eu_doc(
            product,
            findings,
            "EU_GPSR",
            "has_gpsr_docs",
            "GPSR 技术文件与风险评估",
            "high",
            "补齐 GPSR 风险评估、追溯信息、事故记录与召回流程。",
        )
        self._require_eu_doc(
            product,
            findings,
            "EU_GPSR",
            "has_eu_rep",
            "EU Responsible Person",
            "high",
            "指定欧盟责任人，并在包装/说明书/Listing 标注名称、地址、邮箱。",
        )
        if not product.get("has_ce_mark"):
            self._add(
                product,
                findings,
                market="EU",
                standard_id="EU_LVD",
                requirement="CE 标记与 EU Declaration of Conformity",
                observed="has_ce_mark=false",
                severity="high",
                suggestion="完成适用指令评估并签署 DoC 后再加贴 CE；不得先贴标后补测试。",
            )
        if product.get("wireless") and not product.get("has_red_report"):
            self._add(
                product,
                findings,
                market="EU",
                standard_id="EU_RED",
                requirement="无线模块适用 RED 2014/53/EU",
                observed=f"wireless=true ({product.get('wireless_tech') or '未注明'}) / has_red_report=false",
                severity="high",
                suggestion="补齐 RED 射频、EMC、安全与射频暴露评估；确认模块证书覆盖整机。",
            )
        required_languages = self.manual_languages.get("eu_required", ["en", "de", "fr"])
        languages = {str(x).lower() for x in (product.get("manual_languages") or [])}
        missing_languages = [lang for lang in required_languages if lang.lower() not in languages]
        if missing_languages:
            self._add(
                product,
                findings,
                market="EU",
                standard_id="EU_GPSR",
                requirement="说明书/安全信息应覆盖目标市场语言",
                observed=f"manual_languages={sorted(languages)}，缺少 {missing_languages}",
                severity="medium",
                suggestion="补充目标市场官方语言说明书与安全警告，DE/FR 需母语审校。",
            )
        plug = str(product.get("plug_type") or "")
        if not any(token in plug.upper() for token in ("EU", "TYPE C", "TYPE F", "SCHUKO")):
            self._add(
                product,
                findings,
                market="EU",
                standard_id="EU_LVD",
                requirement="欧盟市场插头/接口应与 230V 制式匹配",
                observed=f"plug_type={plug or '未填写'}",
                severity="high",
                suggestion="提供 EU Type C/F 插头版本，或在 Listing 明确该 SKU 不适用 EU。",
            )
        return findings

    def _check_us(self, product: dict[str, Any]) -> list[Finding]:
        findings: list[Finding] = []
        voltage = _num(product.get("output_voltage_v"))
        frequency = _num(product.get("output_frequency_hz"))
        low, high = [float(x) for x in self.thresholds.get("us_voltage_v", [100, 130])]
        us_freq = float(self.thresholds.get("us_frequency_hz", 60))
        if not low <= voltage <= high:
            self._add(
                product,
                findings,
                market="US",
                standard_id="US_FCC15B",
                requirement="美国市场输出额定电压应为 100-130V",
                observed=f"当前输出 {voltage:g}V / {frequency:g}Hz",
                severity="high",
                suggestion="为美国版本使用 120V/60Hz 逆变器；同步更新铭牌、包装和 Listing。",
            )
        if frequency != us_freq:
            self._add(
                product,
                findings,
                market="US",
                standard_id="US_FCC15B",
                requirement="美国市场输出频率应为 60Hz",
                observed=f"当前频率 {frequency:g}Hz",
                severity="high",
                suggestion="更换 US 版本逆变器/主板并重新测试。",
            )
        self._require_us_doc(
            product,
            findings,
            "US_FCC15B",
            "has_fcc_part15b",
            "FCC Part 15 Subpart B 测试报告",
            "high",
            "补齐 Part 15B 发射测试，并准备 SDoC。",
        )
        self._require_us_doc(
            product,
            findings,
            "US_FCC15B",
            "has_fcc_sdoc",
            "FCC SDoC 符合性声明",
            "high",
            "由美国责任方签署 SDoC，保留测试报告与标签。",
        )
        if not product.get("has_fcc_label"):
            self._add(
                product,
                findings,
                market="US",
                standard_id="US_FCC15B",
                requirement="FCC 合规标签/用户手册声明",
                observed="has_fcc_label=false",
                severity="medium",
                suggestion="按规则添加 FCC 合规声明与必要标签。",
            )
        if product.get("wireless") and not product.get("has_fcc_id"):
            self._add(
                product,
                findings,
                market="US",
                standard_id="US_FCC15C",
                requirement="无线模块需 FCC ID / Part 15C 认证",
                observed=f"wireless=true ({product.get('wireless_tech') or '未注明'}) / has_fcc_id=false",
                severity="high",
                suggestion="完成 FCC Part 15C 认证并取得 FCC ID；确认模块认证覆盖整机。",
            )
        if not product.get("has_ul2743"):
            self._add(
                product,
                findings,
                market="US",
                standard_id="US_UL2743",
                requirement="美国零售市场通常要求 UL 2743 或 NRTL 认可的安全方案",
                observed="has_ul2743=false",
                severity="medium",
                suggestion="安排 UL 2743 测试或提供 NRTL 认可的等效安全报告。",
            )
        if not product.get("has_prop65_warning"):
            self._add(
                product,
                findings,
                market="US",
                standard_id="US_PROP65",
                requirement="加州 Prop 65 化学品暴露评估/警示",
                observed="has_prop65_warning=false",
                severity="medium",
                suggestion="完成材料评估，必要时添加 Prop 65 警示标签。",
            )
        plug = str(product.get("plug_type") or "")
        if not any(token in plug.upper() for token in ("US", "TYPE A", "TYPE B")):
            self._add(
                product,
                findings,
                market="US",
                standard_id="US_FCC15B",
                requirement="美国市场插头/接口应与 120V 制式匹配",
                observed=f"plug_type={plug or '未填写'}",
                severity="high",
                suggestion="提供 US Type A/B 插头版本，或在 Listing 明确该 SKU 不适用 US。",
            )
        return findings

    def _check_general(self, product: dict[str, Any]) -> list[Finding]:
        findings: list[Finding] = []
        battery_wh = _num(product.get("battery_wh"))
        if battery_wh <= 0:
            self._add(
                product,
                findings,
                market="GLOBAL",
                standard_id="UN383",
                requirement="锂电池产品必须明确容量与 UN38.3 状态",
                observed="battery_wh<=0 或未填写",
                severity="high",
                suggestion="补齐电池容量、电芯规格与 UN38.3 测试摘要，再评估运输与上架。",
            )
        elif not product.get("has_un383"):
            self._add(
                product,
                findings,
                market="GLOBAL",
                standard_id="UN383",
                requirement="锂电池运输需 UN38.3 测试摘要",
                observed="has_un383=false",
                severity="high",
                suggestion="取得 UN38.3 测试摘要、MSDS 与合规包装方案后再安排运输。",
            )
        if battery_wh > float(self.thresholds.get("battery_air_advisory_wh", 100)):
            self._add(
                product,
                findings,
                market="GLOBAL",
                standard_id="IATA",
                requirement=">100Wh 锂电池空运通常受限制/需批准",
                observed=f"battery_wh={battery_wh:g}",
                severity="low",
                suggestion="核对 IATA DGR 包装说明、航空运输鉴定与货代要求；必要时改海运/陆运。",
            )
        if not product.get("ip_rating"):
            self._add(
                product,
                findings,
                market="GLOBAL",
                standard_id="EU_LVD",
                requirement="户外使用场景应明确 IP 防护等级",
                observed="ip_rating 未填写",
                severity="medium",
                suggestion="补充 IP 等级与测试依据；未测试前不要在 Listing 使用防水/防尘表述。",
            )
        if _num(product.get("rated_power_w")) < _num(product.get("output_w")):
            self._add(
                product,
                findings,
                market="GLOBAL",
                standard_id="EU_LVD",
                requirement="额定功率不得低于对外宣称的 AC 输出功率",
                observed=f"rated_power_w={product.get('rated_power_w')} / output_w={product.get('output_w')}",
                severity="high",
                suggestion="核对铭牌、说明书与测试报告的额定功率，统一口径。",
            )
        return findings

    def _require_eu_doc(
        self,
        product: dict[str, Any],
        findings: list[Finding],
        standard_id: str,
        field: str,
        requirement: str,
        severity: str,
        suggestion: str,
    ) -> None:
        if product.get(field):
            return
        self._add(
            product,
            findings,
            market="EU",
            standard_id=standard_id,
            requirement=requirement,
            observed=f"{field}=false",
            severity=severity,
            suggestion=suggestion,
        )

    def _require_us_doc(
        self,
        product: dict[str, Any],
        findings: list[Finding],
        standard_id: str,
        field: str,
        requirement: str,
        severity: str,
        suggestion: str,
    ) -> None:
        if product.get(field):
            return
        self._add(
            product,
            findings,
            market="US",
            standard_id=standard_id,
            requirement=requirement,
            observed=f"{field}=false",
            severity=severity,
            suggestion=suggestion,
        )

    def _summary_row(
        self,
        product: dict[str, Any],
        findings: list[Finding],
        markets: list[str],
    ) -> dict[str, Any]:
        counts = Counter(item.severity for item in findings)
        score = 100.0
        for severity, count in counts.items():
            score -= count * self.weights.get(severity, 10)
        score = max(0.0, min(100.0, score))
        if counts["high"]:
            status = "FAIL"
        elif counts["medium"]:
            status = "CONDITIONAL"
        else:
            status = "PASS"
        market_status: dict[str, str] = {}
        for market in markets:
            market_findings = [item for item in findings if item.market == market]
            if any(item.severity == "high" for item in market_findings):
                market_status[market] = "FAIL"
            elif any(item.severity == "medium" for item in market_findings):
                market_status[market] = "CONDITIONAL"
            else:
                market_status[market] = "PASS"
        top = next((item for item in findings if item.severity in {"high", "medium"}), None)
        return {
            "sku": product["sku"],
            "model": product.get("model") or "",
            "name": product.get("name") or "",
            "platform": product.get("platform_label") or "",
            "target_markets": "/".join(markets),
            "output_voltage_v": product.get("output_voltage_v"),
            "output_frequency_hz": product.get("output_frequency_hz"),
            "rated_power_w": product.get("rated_power_w"),
            "battery_wh": product.get("battery_wh"),
            "battery_chemistry": product.get("battery_chemistry"),
            "wireless": bool(product.get("wireless")),
            "high": counts["high"],
            "medium": counts["medium"],
            "low": counts["low"],
            "risk_score": round(score, 1),
            "status": status,
            "us_status": market_status.get("US", "-"),
            "eu_status": market_status.get("EU", "-"),
            "top_risk": top.requirement if top else "未发现规则内高风险",
            "top_suggestion": top.suggestion if top else "保持文件更新并做年度复核",
        }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def export_excel(
    rows: list[dict[str, Any]],
    findings: list[Finding],
    products: list[dict[str, Any]],
    rules: Rules,
    output: Path,
    *,
    run_id: str,
    generated_at: str,
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ComplianceError("生成 Excel 需要 openpyxl") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "合规汇总"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    def meta(sheet: Any, title: str, columns: int = 12) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        sheet["A1"] = title
        sheet["A1"].font = Font(bold=True, size=16, color="1F4E78")
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
        sheet["A2"] = f"运行：{run_id} ｜ 生成：{generated_at} ｜ 规则初筛，不构成法律/认证意见"
        sheet["A2"].font = Font(size=9, color="595959")
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=columns)
        sheet["A3"] = "范围：CE（LVD/EMC/RoHS/WEEE/Battery/GPSR/RED）与 FCC（Part 15B/15C）；数据为模拟产品参数。"
        sheet["A3"].font = Font(size=9, color="7F6000")
        sheet["A3"].fill = PatternFill("solid", fgColor="FFF2CC")
        sheet.row_dimensions[1].height = 26
        sheet.row_dimensions[2].height = 20
        sheet.row_dimensions[3].height = 22

    # 合规汇总
    meta(ws, "欧美储能产品合规风险汇总", 15)
    headers = [
        "SKU", "型号", "商品", "平台", "目标市场", "输出电压(V)", "频率(Hz)", "额定功率(W)",
        "电池(Wh)", "高风险", "中风险", "低风险", "风险分", "状态", "美国", "欧盟", "Top风险",
    ]
    header_row = 5
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_index, item in enumerate(rows, start=header_row + 1):
        values = [
            item["sku"], item["model"], item["name"], item["platform"], item["target_markets"],
            item["output_voltage_v"], item["output_frequency_hz"], item["rated_power_w"],
            item["battery_wh"], item["high"], item["medium"], item["low"], item["risk_score"],
            item["status"], item["us_status"], item["eu_status"], item["top_risk"],
        ]
        for col_index, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col_index in {3, 17})
            if col_index == 13:
                cell.number_format = "0.0"
        status_cell = ws.cell(row=row_index, column=14)
        color = {"FAIL": "F4CCCC", "CONDITIONAL": "FFF2CC", "PASS": "D9EAD3"}.get(str(item["status"]), "FFFFFF")
        status_cell.fill = PatternFill("solid", fgColor=color)
        status_cell.font = Font(bold=True)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A{header_row}:Q{header_row + len(rows)}"
    widths = [20, 18, 30, 11, 12, 12, 10, 11, 10, 9, 9, 9, 9, 12, 10, 10, 44]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.sheet_view.showGridLines = False
    if rows:
        ws.conditional_formatting.add(
            f"M{header_row + 1}:M{header_row + len(rows)}",
            ColorScaleRule(
                start_type="num", start_value=0, start_color="F8696B",
                mid_type="num", mid_value=60, mid_color="FFEB84",
                end_type="num", end_value=100, end_color="63BE7B",
            ),
        )

    # 不符合项
    ws2 = wb.create_sheet("不符合项")
    meta(ws2, "不符合项与修改建议", 12)
    f_headers = ["ID", "SKU", "型号/商品", "市场", "标准", "要求", "观察值", "严重级别", "状态", "修改建议", "需要的证据"]
    for index, header in enumerate(f_headers, start=1):
        cell = ws2.cell(row=5, column=index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_index, finding in enumerate(findings, start=6):
        values = [
            getattr(finding, "finding_id", ""), finding.sku, finding.product_name, finding.market,
            finding.standard_label, finding.requirement, finding.observed, finding.severity,
            finding.status, finding.suggestion, finding.evidence,
        ]
        for col_index, value in enumerate(values, start=1):
            cell = ws2.cell(row=row_index, column=col_index, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col_index in {4, 5, 6, 7, 10, 11})
        sev_col = {"high": "F4CCCC", "medium": "FFF2CC", "low": "DDEBF7"}.get(finding.severity, "FFFFFF")
        ws2.cell(row=row_index, column=8).fill = PatternFill("solid", fgColor=sev_col)
    ws2.freeze_panes = "A6"
    ws2.auto_filter.ref = f"A5:K{5 + len(findings)}"
    for index, width in enumerate([8, 20, 28, 8, 32, 42, 34, 10, 10, 48, 48], start=1):
        ws2.column_dimensions[get_column_letter(index)].width = width
    ws2.sheet_view.showGridLines = False

    # 产品参数
    ws3 = wb.create_sheet("产品参数")
    meta(ws3, "模拟产品合规主数据", 12)
    p_headers = [
        "SKU", "型号", "目标市场", "输出电压(V)", "频率(Hz)", "额定功率(W)", "电池(Wh)",
        "电芯", "无线", "无线技术", "IP等级", "插头", "说明书语言", "CE", "LVD", "EMC",
        "RoHS", "WEEE", "Battery", "GPSR", "EU责任人", "RED", "UN38.3", "UL2743",
        "FCC 15B", "SDoC", "FCC ID", "FCC标签", "Prop65",
    ]
    for index, header in enumerate(p_headers, start=1):
        cell = ws3.cell(row=5, column=index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_index, product in enumerate(products, start=6):
        values = [
            product.get("sku"), product.get("model"), "/".join(product.get("target_markets", [])),
            product.get("output_voltage_v"), product.get("output_frequency_hz"), product.get("rated_power_w"),
            product.get("battery_wh"), product.get("battery_chemistry"), "是" if product.get("wireless") else "否",
            product.get("wireless_tech"), product.get("ip_rating"), product.get("plug_type"),
            "/".join(product.get("manual_languages", [])),
            *["是" if product.get(key) else "否" for key in (
                "has_ce_mark", "has_lvd_report", "has_emc_report", "has_rohs_report",
                "has_weee_registration", "has_battery_regulation_doc", "has_gpsr_docs",
                "has_eu_rep", "has_red_report", "has_un383", "has_ul2743",
                "has_fcc_part15b", "has_fcc_sdoc", "has_fcc_id", "has_fcc_label",
                "has_prop65_warning",
            )],
        ]
        for col_index, value in enumerate(values, start=1):
            cell = ws3.cell(row=row_index, column=col_index, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col_index in {2, 3, 13})
    ws3.freeze_panes = "A6"
    ws3.auto_filter.ref = f"A5:AC{5 + len(products)}"
    for index in range(1, len(p_headers) + 1):
        ws3.column_dimensions[get_column_letter(index)].width = 12
    for index in (1, 2, 3, 13):
        ws3.column_dimensions[get_column_letter(index)].width = 22

    # 标准依据
    ws4 = wb.create_sheet("标准依据")
    meta(ws4, "公开标准与证据清单", 8)
    s_headers = ["ID", "标准/法规", "市场", "适用范围", "需要的证据", "来源"]
    for index, header in enumerate(s_headers, start=1):
        cell = ws4.cell(row=5, column=index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row_index, standard in enumerate(rules.standards.values(), start=6):
        values = [
            standard.get("id"), standard.get("label"), standard.get("market"),
            standard.get("scope"), standard.get("evidence"), standard.get("source"),
        ]
        for col_index, value in enumerate(values, start=1):
            cell = ws4.cell(row=row_index, column=col_index, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws4.freeze_panes = "A6"
    for index, width in enumerate([18, 40, 10, 52, 58, 38], start=1):
        ws4.column_dimensions[get_column_letter(index)].width = width
    ws4.sheet_view.showGridLines = False

    wb.properties.title = "欧美储能产品合规风险筛查"
    wb.properties.creator = "Outdoor ESS AI compliance_check.py"
    wb.properties.description = "规则初筛，不构成法律或认证意见。"
    utils.ensure_dir(output.parent)
    wb.save(output)
    logger.info("合规 Excel 已生成: %s", output)
    return output


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="compliance_check",
        description="欧美储能产品 CE/FCC 合规风险筛查",
    )
    parser.add_argument("--selection", default=str(DEFAULT_SELECTION), help="选品表")
    parser.add_argument("--master", default=str(DEFAULT_MASTER), help="产品合规主数据")
    parser.add_argument("--config", default=str(DEFAULT_RULES), help="合规规则")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Excel 输出")
    parser.add_argument("--json", default=str(DEFAULT_JSON), help="JSON 摘要")
    parser.add_argument("--strict", action="store_true", help="存在高风险时退出码 2")
    parser.add_argument("--log-level", default="INFO")
    return parser


def _configure_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if callable(reconfigure):
            reconfigure(encoding="utf-8", errors="replace")


def main(argv: Iterable[str] | None = None) -> int:
    _configure_stdio()
    args = build_parser().parse_args(list(argv) if argv is not None else None)
    try:
        rules = load_rules(args.config)
        selection = _load_selection(args.selection)
        products = load_master(args.master, selection)
        checker = ComplianceChecker(rules)
        rows, findings = checker.check(products)
    except (ComplianceError, utils.UtilError) as exc:
        print(f"✗ 合规筛查失败：{exc}", file=sys.stderr)
        return 1

    output = Path(args.output).expanduser().resolve()
    json_path = Path(args.json).expanduser().resolve()
    utils.ensure_dir(output.parent)
    utils.setup_logging(
        level=args.log_level,
        log_file=output.parent / "compliance_check.log",
        console=True,
        force=True,
    )
    run_id = f"compliance_{datetime.now(UTC).strftime('%Y%m%d_%H%M%S')}"
    generated_at = utils.utc_now_iso()
    try:
        export_excel(rows, findings, products, rules, output, run_id=run_id, generated_at=generated_at)
        summary = {
            "run_id": run_id,
            "generated_at": generated_at,
            "notice": "规则初筛，不构成法律、认证或测试结论；认证状态为模拟数据。",
            "products": len(products),
            "finding_count": len(findings),
            "severity": {
                "high": sum(1 for item in findings if item.severity == "high"),
                "medium": sum(1 for item in findings if item.severity == "medium"),
                "low": sum(1 for item in findings if item.severity == "low"),
            },
            "status": {
                key: sum(1 for item in rows if item["status"] == key)
                for key in ("PASS", "CONDITIONAL", "FAIL")
            },
            "rows": rows,
            "findings": [item.to_dict() for item in findings],
            "outputs": {"excel": str(output), "json": str(json_path)},
        }
        utils.write_json(json_path, summary)
    except (ComplianceError, utils.UtilError, OSError) as exc:
        logger.error("导出失败: %s", exc)
        print(f"✗ 导出失败：{exc}", file=sys.stderr)
        return 1

    print("")
    print("=" * 72)
    print(f"✓ 合规筛查完成：{run_id}")
    print(f"  产品数  : {len(products)}")
    print(
        f"  风险项  : 高 {summary['severity']['high']} / 中 {summary['severity']['medium']} "
        f"/ 低 {summary['severity']['low']}"
    )
    print(
        f"  状态    : PASS {summary['status']['PASS']} / CONDITIONAL "
        f"{summary['status']['CONDITIONAL']} / FAIL {summary['status']['FAIL']}"
    )
    print(f"  Excel   : {output}")
    print(f"  JSON    : {json_path}")
    print("=" * 72)
    if args.strict and summary["severity"]["high"]:
        return 2
    return 0


if __name__ == "__main__":
    from error_guard import run_cli

    raise SystemExit(run_cli(main, "compliance_check"))
