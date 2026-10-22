"""合规筛查与经营报表：公式、结构和 Excel 输出测试。"""

from __future__ import annotations

from pathlib import Path

from compliance_check import (
    DEFAULT_MASTER,
    DEFAULT_RULES,
    DEFAULT_SELECTION,
    ComplianceChecker,
    _load_selection,
    export_excel,
    load_master,
    load_rules,
)
from report_generator import (
    OperationsSimulator,
    load_config,
    load_products,
    write_ad_workbook,
    write_inventory_workbook,
    write_profit_workbook,
)


def test_compliance_rules_and_scores(workspace_tmp: Path) -> None:
    rules = load_rules(DEFAULT_RULES)
    products = load_master(DEFAULT_MASTER, _load_selection(DEFAULT_SELECTION))
    rows, findings = ComplianceChecker(rules).check(products)

    assert len(products) == 9
    assert len(rows) == 9
    assert findings
    assert any(item.standard_id == "EU_LVD" for item in findings)
    assert any(item.standard_id == "US_FCC15B" for item in findings)

    by_sku = {row["sku"]: row for row in rows}
    assert by_sku["1688-600000000005"]["eu_status"] == "PASS"
    assert by_sku["AMZ-B0DEMO0003"]["us_status"] == "PASS"
    assert by_sku["1688-600000000001"]["status"] == "FAIL"

    for row in rows:
        expected = max(
            0.0,
            100.0 - 25.0 * row["high"] - 10.0 * row["medium"] - 5.0 * row["low"],
        )
        assert abs(row["risk_score"] - round(expected, 1)) <= 0.05
        assert row["status"] in {"PASS", "CONDITIONAL", "FAIL"}

    output = workspace_tmp / "compliance.xlsx"
    export_excel(rows, findings, products, rules, output, run_id="test", generated_at="2026-10-21T00:00:00+00:00")
    import openpyxl

    wb = openpyxl.load_workbook(output, data_only=True)
    assert wb.sheetnames == ["合规汇总", "不符合项", "产品参数", "标准依据"]
    assert wb["合规汇总"].max_row == 6 + len(rows) - 1
    assert wb["不符合项"].max_row == 6 + len(findings) - 1
    assert [cell.value for cell in wb["合规汇总"][5]][:4] == ["SKU", "型号", "商品", "平台"]


def test_operations_formulas_and_excel(workspace_tmp: Path) -> None:
    config = load_config()
    products = load_products(config.resolve("input.selection"))
    data = OperationsSimulator(config, products).run()
    ad_rows = data["ad_rows"]
    metrics = data["metrics"]

    assert len(ad_rows) == 270
    assert len(metrics) == 9
    assert data["summary"]["units"] == sum(item["units"] for item in metrics)

    for row in ad_rows:
        assert abs(row["ctr"] - row["clicks"] / row["impressions"]) <= 1e-6
        assert abs(row["cvr"] - row["ad_orders"] / row["clicks"]) <= 1e-6
        assert abs(row["acos"] - row["spend"] / row["ad_sales"]) <= 1e-5
        assert abs(row["roas"] - row["ad_sales"] / row["spend"]) <= 0.01
        assert abs(
            row["ad_profit"]
            - round(row["ad_sales"] * row["gross_margin_pct"] / 100 - row["spend"], 2)
        ) <= 0.01

    for item in metrics:
        assert abs(item["revenue"] - item["units"] * item["price"]) <= 0.02
        assert abs(item["gross_profit"] - (item["revenue"] - item["cogs"])) <= 0.02
        assert abs(item["net_profit"] - (item["revenue"] - item["total_cost"])) <= 0.02
        assert abs(item["roi"] - item["net_profit"] / item["total_cost"]) <= 1e-9
        assert abs(item["net_margin"] - item["net_profit"] / item["revenue"]) <= 1e-9
        assert abs(
            item["inventory_turnover_days"]
            - round(item["avg_inventory_units"] / item["avg_daily_units"], 1)
        ) <= 0.06
        assert abs(item["days_cover"] - round((item["on_hand"] + item["in_transit"]) / item["avg_daily_units"], 1)) <= 0.06
        assert abs(item["reorder_point"] - round(item["avg_daily_units"] * (item["lead_time_days"] + item["safety_days"]), 1)) <= 0.06
        assert item["suggested_order_qty"] >= 0
        if item["days_cover"] < item["lead_time_days"]:
            assert item["stock_status"] == "紧急补货"
        elif item["days_cover"] < item["lead_time_days"] + item["safety_days"]:
            assert item["stock_status"] == "关注"
        elif item["days_cover"] > 90:
            assert item["stock_status"] == "滞销"
        else:
            assert item["stock_status"] == "健康"

    summary = data["summary"]
    total_cost = sum(item["total_cost"] for item in metrics)
    assert abs(summary["roi"] - summary["net_profit"] / total_cost * 100) <= 0.02
    assert abs(summary["net_margin"] - summary["net_profit"] / summary["revenue"] * 100) <= 0.02

    ad_path = workspace_tmp / "ad.xlsx"
    inventory_path = workspace_tmp / "inventory.xlsx"
    profit_path = workspace_tmp / "profit.xlsx"
    write_ad_workbook(ad_path, ad_rows, summary, "test", data["period"])
    write_inventory_workbook(inventory_path, metrics, "test", data["period"])
    write_profit_workbook(profit_path, metrics, summary, "test", data["period"])
    import openpyxl

    wb_ad = openpyxl.load_workbook(ad_path, data_only=True)
    assert wb_ad.sheetnames == ["广告日报", "指标说明"]
    assert wb_ad["广告日报"].max_row == 277
    assert wb_ad["广告日报"].cell(row=7, column=1).value is not None

    wb_inv = openpyxl.load_workbook(inventory_path, data_only=True)
    assert wb_inv.sheetnames == ["库存预警", "指标说明"]
    assert wb_inv["库存预警"].max_row == 15

    wb_profit = openpyxl.load_workbook(profit_path, data_only=True)
    assert wb_profit.sheetnames == ["经营概览", "月度利润", "指标说明"]
    assert wb_profit["月度利润"].max_row == 16
    assert wb_profit["经营概览"]["A1"].value.startswith("月度经营概览")
