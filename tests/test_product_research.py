"""选品模块：清洗、指标自洽性与 Excel 输出测试。"""

from __future__ import annotations

from pathlib import Path

import utils
from product_research import (
    Alibaba1688Adapter,
    AliExpressAdapter,
    AmazonAdapter,
    ProductCleaner,
    ProductResearchEngine,
    PublicFetcher,
    _extract_spec,
    apply_overrides,
    load_research_config,
)


def test_extract_spec_from_title() -> None:
    title = "300Wh Portable Power Station 300W 4.2kg"
    assert _extract_spec(title, "capacity_wh") == 300
    assert _extract_spec(title, "output_w") == 300
    assert _extract_spec(title, "weight_kg") == 4.2


def test_public_page_adapters_parse_minimal_cards() -> None:
    config = load_research_config()
    fetcher = PublicFetcher(config)
    platforms = dict(config.platform_items())

    amazon_html = """
    <div data-component-type="s-search-result" data-asin="B0TEST0001">
      <h2><a href="/dp/B0TEST0001"><span>Demo 300Wh Portable Power Station 300W</span></a></h2>
      <span class="a-price"><span class="a-offscreen">$299.00</span></span>
      <span class="a-icon-alt">4.5 out of 5 stars</span>
      <a href="/product-reviews/B0TEST0001"><span>1,234</span></a>
    </div>
    """
    amazon = AmazonAdapter("amazon", platforms["amazon"], fetcher).parse(
        amazon_html, "https://www.amazon.com/s?k=test"
    )
    assert amazon[0]["product_id"] == "B0TEST0001"
    assert amazon[0]["price_local"] == 299.0
    assert amazon[0]["review_count"] == 1234

    aliexpress_html = """
    <div class="search-item-card-wrapper-gallery">
      <a href="https://www.aliexpress.com/item/1005009999999.html">
        <h3>Demo 200Wh Portable Power Station 200W</h3>
        <div class="seo-optimized-price">US $139.00</div>
      </a>
    </div>
    """
    ali = AliExpressAdapter("aliexpress", platforms["aliexpress"], fetcher).parse(
        aliexpress_html, "https://www.aliexpress.com/w/wholesale-test.html"
    )
    assert ali[0]["product_id"] == "1005009999999"
    assert ali[0]["price_local"] == 139.0

    alibaba_html = """
    <div data-offer-id="600000009999">
      <a href="https://detail.1688.com/offer/600000009999.html">
        <div class="title-text">户外电源 300Wh 露营便携式储能电源</div>
        <div class="price">¥780.00</div>
      </a>
    </div>
    """
    alibaba = Alibaba1688Adapter("alibaba_1688", platforms["alibaba_1688"], fetcher).parse(
        alibaba_html, "https://s.1688.com/selloffer/offer_search.htm?keywords=test"
    )
    assert alibaba[0]["product_id"] == "600000009999"
    assert alibaba[0]["price_local"] == 780.0


def test_cleaner_filters_and_dedupes() -> None:
    config = load_research_config()
    cleaner = ProductCleaner(config)
    fixture = utils.read_json(config.resolve("fixture.path"))
    records = fixture["records"]
    duplicated = dict(records[0])
    duplicated["product_id"] = "DUPLICATE-ID"
    rows, stats = cleaner.clean([*records, duplicated])
    assert stats["clean_records"] == len(rows)
    assert all(100 <= row["capacity_wh"] <= 500 for row in rows)
    assert stats["dropped_by_reason"].get("capacity_out_of_range", 0) >= 4
    assert stats["dropped_by_reason"].get("duplicate_title", 0) >= 1


def test_full_research_run_and_metric_consistency(workspace_tmp: Path) -> None:
    config = load_research_config()
    config = apply_overrides(config, output_dir=str(workspace_tmp / "research"))
    summary = ProductResearchEngine(config).run(mode="offline")

    assert summary["raw_records"] == 26
    assert summary["cleaned_records"] == 20
    assert summary["high_potential_records"] >= 5
    assert summary["outputs"]["summary"].endswith("research_summary.json")

    rows = utils.read_csv(config.root / summary["outputs"]["csv"])
    assert len(rows) == summary["cleaned_records"]
    for row in rows:
        capacity = float(row["capacity_wh"])
        weight = float(row["weight_kg"])
        retail = float(row["retail_price_usd"])
        landed = float(row["landed_cost_usd"])
        fee = float(row["platform_fee_usd"])
        fulfillment = float(row["fulfillment_cost_usd"])
        gross = float(row["gross_profit_usd"])
        margin = float(row["gross_margin_pct"])
        market = float(row["market_capacity_usd"])
        monthly_units = int(float(row["monthly_units"]))
        competition = float(row["competition_score"])

        assert 100 <= capacity <= 500
        assert weight > 0
        assert retail > 0
        assert abs(gross - round(retail - landed - fee - fulfillment, 2)) <= 0.02
        assert abs(margin - round(gross / retail * 100, 2)) <= 0.05
        assert abs(market - round(monthly_units * retail, 2)) <= 0.02
        assert 0 <= competition <= 100

    passing = [row for row in rows if row["gate_status"] == "通过"]
    assert len(passing) == summary["high_potential_records"]
    for row in passing:
        assert float(row["gross_margin_pct"]) >= summary["filters"]["min_gross_margin_pct"]
        assert float(row["market_capacity_usd"]) >= summary["filters"]["min_market_capacity_usd"]
        assert float(row["competition_score"]) <= summary["filters"]["max_competition_score"]
        assert float(row["opportunity_score"]) >= summary["filters"]["min_opportunity_score"]


def test_excel_structure_and_headers(workspace_tmp: Path) -> None:
    config = load_research_config()
    config = apply_overrides(config, output_dir=str(workspace_tmp / "excel"))
    summary = ProductResearchEngine(config).run(mode="offline")
    workbook_path = config.root / summary["outputs"]["excel"]

    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    assert wb.sheetnames == ["高潜力选品", "全部清洗数据", "指标说明", "筛选参数", "数据质量"]

    high_ws = wb["高潜力选品"]
    headers = [cell.value for cell in high_ws[6]]
    assert headers[:4] == ["排名", "平台", "商品标题", "品牌/供应商"]
    assert "毛利率(%)" in headers
    assert "市场容量(USD/月)" in headers
    assert "竞争度(越低越好)" in headers
    assert "综合机会分" in headers
    assert high_ws.max_row == 6 + summary["high_potential_records"]
    assert high_ws["A1"].value.startswith("户外露营轻量化储能")

    column = {name: index + 1 for index, name in enumerate(headers)}
    for row_index in range(7, high_ws.max_row + 1):
        retail = high_ws.cell(row=row_index, column=column["零售基准价(USD)"]).value
        landed = high_ws.cell(row=row_index, column=column["到岸成本(USD)"]).value
        fee = high_ws.cell(row=row_index, column=column["平台佣金(USD)"]).value
        fulfillment = high_ws.cell(row=row_index, column=column["履约成本(USD)"]).value
        gross = high_ws.cell(row=row_index, column=column["单件毛利(USD)"]).value
        margin = high_ws.cell(row=row_index, column=column["毛利率(%)"]).value
        market = high_ws.cell(row=row_index, column=column["市场容量(USD/月)"]).value
        units = high_ws.cell(row=row_index, column=column["预估月销量"]).value
        assert abs(gross - round(retail - landed - fee - fulfillment, 2)) <= 0.02
        assert abs(margin - round(gross / retail * 100, 2)) <= 0.05
        assert abs(market - round(units * retail, 2)) <= 0.02
        assert high_ws.cell(row=row_index, column=column["门槛状态"]).value == "通过"

    all_ws = wb["全部清洗数据"]
    assert all_ws.max_row == 1 + summary["cleaned_records"]
    assert all_ws.cell(row=2, column=1).value is not None

    assert wb["指标说明"].max_row >= 8
    assert wb["筛选参数"].max_row > 10
    assert wb["数据质量"].max_row > 5
