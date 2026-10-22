#!/usr/bin/env python3
"""经营报表自动化：广告日报、库存预警、月度利润。

所有经营数据为固定种子模拟，不接入真实店铺、广告或库存后台。
"""

from __future__ import annotations

import argparse
import math
import random
import re
import sys
from collections.abc import Iterable
from datetime import date, timedelta
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parent
VENDOR_DIR = PROJECT_ROOT / ".vendor"
if VENDOR_DIR.is_dir() and str(VENDOR_DIR) not in sys.path:
    sys.path.insert(0, str(VENDOR_DIR))

import utils  # noqa: E402

logger = utils.get_logger("report_generator")
DEFAULT_CONFIG = PROJECT_ROOT / "config" / "report_generator.yaml"
DEFAULT_INPUT = PROJECT_ROOT / "outputs" / "research" / "high_potential_products.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "outputs" / "operations"


class ReportError(RuntimeError):
    """报表生成异常。"""


# ---------------------------------------------------------------------------
# 配置与选品表
# ---------------------------------------------------------------------------
class ReportConfig:
    def __init__(self, root: Path, path: Path, data: dict[str, Any]) -> None:
        self.root = root
        self.path = path
        self.data = data

    def get(self, dotted: str, default: Any = None) -> Any:
        return utils.deep_get(self.data, dotted, default)

    def resolve(self, dotted: str) -> Path:
        raw = self.get(dotted)
        if not raw:
            raise ReportError(f"缺少路径配置: {dotted}")
        path = Path(str(raw)).expanduser()
        return path if path.is_absolute() else (self.root / path).resolve()


def load_config(path: str | Path | None = None, *, root: str | Path | None = None) -> ReportConfig:
    project_root = Path(root).expanduser().resolve() if root else PROJECT_ROOT
    config_path = Path(path).expanduser() if path else project_root / "config" / "report_generator.yaml"
    if not config_path.is_absolute():
        config_path = (project_root / config_path).resolve()
    data = utils.read_yaml(config_path)
    if not isinstance(data, dict):
        raise ReportError(f"配置格式错误: {config_path}")
    if not data.get("input") or not data.get("output"):
        raise ReportError("配置缺少 input/output")
    return ReportConfig(project_root, config_path, data)


def _num(value: Any, default: float = 0.0) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", str(value or ""))
    return float(match.group(0).replace(",", "")) if match else default


def load_products(path: str | Path, sheet: str = "高潜力选品") -> list[dict[str, Any]]:
    source = Path(path).expanduser().resolve()
    if not source.exists():
        raise ReportError(f"选品表不存在: {source}")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ReportError("读取选品表需要 openpyxl") from exc
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    headers: list[str] = []
    header_row = 1
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True), start=1):
        values = [str(cell or "").strip() for cell in row]
        if "SKU" in values or "sku" in values:
            headers = values
            header_row = row_index
            break
    products: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        record = dict(zip(headers, row, strict=False))
        sku = str(record.get("SKU") or record.get("sku") or "").strip()
        if not sku:
            continue
        price = _num(record.get("零售基准价(USD)") or record.get("retail_price_usd") or record.get("公开价(USD)"))
        landed = _num(record.get("到岸成本(USD)") or record.get("landed_cost_usd"))
        gross = _num(record.get("毛利率(%)") or record.get("gross_margin_pct"))
        if price <= 0 or landed <= 0:
            logger.warning("跳过价格/成本缺失 SKU: %s", sku)
            continue
        if gross <= 0:
            gross = (price - landed) / price * 100.0
        products.append(
            {
                "sku": sku,
                "name": str(record.get("商品标题") or record.get("title") or sku),
                "platform": str(record.get("平台") or record.get("platform_label") or ""),
                "market": str(record.get("目标市场") or record.get("market") or "US"),
                "price": round(price, 2),
                "landed_cost": round(landed, 2),
                "gross_margin_pct": gross,
                "weight_kg": max(0.1, _num(record.get("重量(kg)") or record.get("weight_kg"), 1.0)),
                "capacity_wh": _num(record.get("容量(Wh)") or record.get("capacity_wh")),
                "market_capacity_usd": _num(record.get("市场容量(USD/月)") or record.get("market_capacity_usd")),
                "opportunity_score": _num(record.get("综合机会分") or record.get("opportunity_score")),
                "moq": int(_num(record.get("MOQ") or record.get("moq"), 1)),
            }
        )
    wb.close()
    if not products:
        raise ReportError("选品表中没有可用 SKU")
    return products


# ---------------------------------------------------------------------------
# 经营模拟
# ---------------------------------------------------------------------------
def _weekday_cn(value: date) -> str:
    return ("周一", "周二", "周三", "周四", "周五", "周六", "周日")[value.weekday()]


class OperationsSimulator:
    def __init__(self, config: ReportConfig, products: list[dict[str, Any]]) -> None:
        self.config = config
        self.products = products
        self.sim = config.data.get("simulation", {}) or {}
        self.cost = config.data.get("cost", {}) or {}
        self.seed = int(self.sim.get("seed", 20261021))
        self.days = max(1, int(self.sim.get("days", 30)))
        self.end = date.fromisoformat(str(self.sim.get("period_end") or "2026-09-30"))
        self.start = self.end - timedelta(days=self.days - 1)
        self.rng = random.Random(self.seed)

    def run(self) -> dict[str, Any]:
        ad_rows: list[dict[str, Any]] = []
        metrics: list[dict[str, Any]] = []
        for product in self.products:
            daily = self._daily_rows(product)
            ad_rows.extend(daily)
            metrics.append(self._aggregate(product, daily))
        ad_rows.sort(key=lambda row: (row["date"], row["sku"]))
        metrics.sort(key=lambda row: row["roi"], reverse=True)
        return {
            "ad_rows": ad_rows,
            "metrics": metrics,
            "summary": self._summary(metrics, ad_rows),
            "period": f"{self.start.isoformat()} ~ {self.end.isoformat()}",
            "days": self.days,
            "seed": self.seed,
        }

    def _daily_rows(self, product: dict[str, Any]) -> list[dict[str, Any]]:
        price = float(product["price"])
        gross_margin = max(0.05, min(0.75, float(product["gross_margin_pct"]) / 100.0))
        monthly_potential = float(product["market_capacity_usd"]) / price if price else 0
        target_share = float(self.sim.get("target_share", 0.08))
        base_daily = max(1.0, monthly_potential * target_share / self.days)
        ad_low, ad_high = [float(x) for x in self.sim.get("ad_order_share", [0.35, 0.65])]
        acos_low, acos_high = [float(x) for x in self.sim.get("target_acos_range", [0.12, 0.22])]
        ctr_low, ctr_high = [float(x) for x in self.sim.get("ctr_range", [0.004, 0.012])]
        cvr_low, cvr_high = [float(x) for x in self.sim.get("cvr_range", [0.04, 0.10])]
        rows: list[dict[str, Any]] = []
        for offset in range(self.days):
            day = self.start + timedelta(days=offset)
            daily_total = max(1, int(round(base_daily * self.rng.uniform(0.6, 1.4))))
            ad_orders = max(1, min(daily_total, int(round(daily_total * self.rng.uniform(ad_low, ad_high)))))
            organic_orders = max(0, daily_total - ad_orders)
            ad_sales = round(ad_orders * price, 2)
            target_acos = self.rng.uniform(acos_low, acos_high)
            spend = round(ad_sales * target_acos, 2)
            ctr = self.rng.uniform(ctr_low, ctr_high)
            cvr = self.rng.uniform(cvr_low, cvr_high)
            clicks = max(5, int(round(ad_orders / cvr)))
            cpc = spend / clicks if clicks else 0.0
            impressions = max(clicks, int(round(clicks / ctr)))
            ctr_actual = clicks / impressions if impressions else 0.0
            cvr_actual = ad_orders / clicks if clicks else 0.0
            acos = spend / ad_sales if ad_sales else 0.0
            roas = ad_sales / spend if spend else 0.0
            ad_profit = round(ad_sales * gross_margin - spend, 2)
            rows.append(
                {
                    "date": day.isoformat(),
                    "weekday": _weekday_cn(day),
                    "platform": product["platform"],
                    "sku": product["sku"],
                    "name": product["name"],
                    "market": product["market"],
                    "price": price,
                    "impressions": impressions,
                    "clicks": clicks,
                    "ctr": round(ctr_actual, 6),
                    "cpc": round(cpc, 2),
                    "spend": spend,
                    "ad_orders": ad_orders,
                    "cvr": round(cvr_actual, 6),
                    "ad_sales": ad_sales,
                    "acos": round(acos, 6),
                    "roas": round(roas, 4),
                    "ad_profit": ad_profit,
                    "gross_margin_pct": round(gross_margin * 100, 2),
                    "total_orders": daily_total,
                    "organic_orders": organic_orders,
                }
            )
        return rows

    def _aggregate(self, product: dict[str, Any], daily: list[dict[str, Any]]) -> dict[str, Any]:
        price = float(product["price"])
        landed = float(product["landed_cost"])
        weight = float(product["weight_kg"])
        units = sum(int(row["total_orders"]) for row in daily)
        ad_orders = sum(int(row["ad_orders"]) for row in daily)
        ad_sales = round(sum(float(row["ad_sales"]) for row in daily), 2)
        ad_spend = round(sum(float(row["spend"]) for row in daily), 2)
        revenue = round(units * price, 2)
        cogs = round(units * landed, 2)
        fee_rate = float(self.cost.get("platform_fee_rate", 0.15))
        platform_fee = round(revenue * fee_rate, 2)
        fulfillment = round(
            units
            * (
                float(self.cost.get("fulfillment_base_usd", 3.0))
                + weight * float(self.cost.get("fulfillment_per_kg_usd", 1.8))
            ),
            2,
        )
        refunds = round(revenue * float(self.cost.get("return_rate", 0.03)), 2)
        other = round(revenue * float(self.cost.get("other_rate", 0.01)), 2)
        avg_daily = round(units / self.days, 2) if self.days else 0.0
        begin_low, begin_high = [float(x) for x in self.sim.get("inventory_begin_cover_range", [30, 90])]
        onhand_low, onhand_high = [float(x) for x in self.sim.get("inventory_onhand_cover_range", [5, 95])]
        transit_low, transit_high = [float(x) for x in self.sim.get("inventory_intransit_cover_range", [0, 30])]
        begin_inventory = max(1, int(math.ceil(avg_daily * self.rng.uniform(begin_low, begin_high))))
        on_hand = max(0, int(math.ceil(avg_daily * self.rng.uniform(onhand_low, onhand_high))))
        in_transit = max(0, int(math.ceil(avg_daily * self.rng.uniform(transit_low, transit_high))))
        avg_inventory_units = round((begin_inventory + on_hand) / 2.0, 2)
        stock_value = round(on_hand * landed, 2)
        inventory_value = round(avg_inventory_units * landed, 2)
        storage = round(inventory_value * float(self.cost.get("storage_rate_month", 0.01)), 2)
        total_cost = round(cogs + platform_fee + fulfillment + ad_spend + refunds + storage + other, 2)
        gross_profit = round(revenue - cogs, 2)
        net_profit = round(revenue - total_cost, 2)
        roi = net_profit / total_cost if total_cost else 0.0
        net_margin = net_profit / revenue if revenue else 0.0
        gross_margin = gross_profit / revenue if revenue else 0.0
        turnover_days = avg_inventory_units / avg_daily if avg_daily else 0.0
        turnover_times = 365.0 / turnover_days if turnover_days else 0.0
        days_cover = (on_hand + in_transit) / avg_daily if avg_daily else 0.0
        lead_time = int(self.rng.choice(self.sim.get("lead_time_choices", [25, 30, 35, 40])))
        safety_days = int(self.sim.get("safety_days", 14))
        reorder_point = round(avg_daily * (lead_time + safety_days), 1)
        target_cover = int(self.sim.get("target_cover_days", 30))
        target_units = avg_daily * (lead_time + safety_days + target_cover)
        suggested_order = max(0, int(math.ceil(target_units - (on_hand + in_transit))))
        if days_cover < lead_time:
            status = "紧急补货"
            priority = "P1"
        elif days_cover < lead_time + safety_days:
            status = "关注"
            priority = "P2"
        elif days_cover > 90:
            status = "滞销"
            priority = "P3"
        else:
            status = "健康"
            priority = "P4"
        aging_units = (
            max(0, int(math.ceil(on_hand * self.rng.uniform(*self.sim.get("aging_ratio_range", [0, 0.3])))))
            if days_cover > 90
            else 0
        )
        return {
            **product,
            "units": units,
            "ad_orders": ad_orders,
            "ad_sales": ad_sales,
            "ad_spend": ad_spend,
            "revenue": revenue,
            "cogs": cogs,
            "gross_profit": gross_profit,
            "gross_margin": gross_margin,
            "platform_fee": platform_fee,
            "fulfillment": fulfillment,
            "refunds": refunds,
            "storage": storage,
            "other": other,
            "net_profit": net_profit,
            "net_margin": net_margin,
            "roi": roi,
            "total_cost": total_cost,
            "ad_acos": ad_spend / ad_sales if ad_sales else 0.0,
            "ad_roas": ad_sales / ad_spend if ad_spend else 0.0,
            "avg_daily_units": round(avg_daily, 2),
            "begin_inventory": begin_inventory,
            "on_hand": on_hand,
            "in_transit": in_transit,
            "avg_inventory_units": avg_inventory_units,
            "stock_value": stock_value,
            "inventory_value": inventory_value,
            "inventory_turnover_days": round(turnover_days, 1),
            "inventory_turnover_times": round(turnover_times, 2),
            "days_cover": round(days_cover, 1),
            "lead_time_days": lead_time,
            "safety_days": safety_days,
            "reorder_point": reorder_point,
            "suggested_order_qty": suggested_order,
            "aging_units": aging_units,
            "stock_status": status,
            "priority": priority,
        }

    def _summary(self, metrics: list[dict[str, Any]], ad_rows: list[dict[str, Any]]) -> dict[str, Any]:
        total_revenue = sum(item["revenue"] for item in metrics)
        total_cost = sum(item["total_cost"] for item in metrics)
        total_net = sum(item["net_profit"] for item in metrics)
        total_ad_spend = sum(item["ad_spend"] for item in metrics)
        total_ad_sales = sum(item["ad_sales"] for item in metrics)
        total_units = sum(item["units"] for item in metrics)
        inventory_alerts = [item for item in metrics if item["stock_status"] != "健康"]
        return {
            "ad_rows": len(ad_rows),
            "sku_count": len(metrics),
            "units": total_units,
            "revenue": round(total_revenue, 2),
            "cogs": round(sum(item["cogs"] for item in metrics), 2),
            "ad_spend": round(total_ad_spend, 2),
            "ad_sales": round(total_ad_sales, 2),
            "net_profit": round(total_net, 2),
            "net_margin": round(total_net / total_revenue * 100, 2) if total_revenue else 0.0,
            "roi": round(total_net / total_cost * 100, 2) if total_cost else 0.0,
            "ad_acos": round(total_ad_spend / total_ad_sales * 100, 2) if total_ad_sales else 0.0,
            "ad_roas": round(total_ad_sales / total_ad_spend, 2) if total_ad_spend else 0.0,
            "inventory_alerts": len(inventory_alerts),
            "urgent_alerts": sum(1 for item in metrics if item["stock_status"] == "紧急补货"),
            "slow_moving": sum(1 for item in metrics if item["stock_status"] == "滞销"),
            "top_roi": max(metrics, key=lambda item: item["roi"])["sku"] if metrics else "",
            "lowest_cover": min(metrics, key=lambda item: item["days_cover"])["sku"] if metrics else "",
        }


# ---------------------------------------------------------------------------
# Excel 输出
# ---------------------------------------------------------------------------
def _require_openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ReportError("生成 Excel 需要 openpyxl") from exc
    return openpyxl


def _styles() -> dict[str, Any]:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    return {
        "header_fill": PatternFill("solid", fgColor="1F4E78"),
        "header_font": Font(color="FFFFFF", bold=True, size=10),
        "border": Border(
            left=Side(style="thin", color="D9E2F3"),
            right=Side(style="thin", color="D9E2F3"),
            top=Side(style="thin", color="D9E2F3"),
            bottom=Side(style="thin", color="D9E2F3"),
        ),
        "alignment": Alignment(vertical="top", wrap_text=True),
    }


def _meta(ws: Any, title: str, subtitle: str, note: str, columns: int) -> None:
    from openpyxl.styles import Font, PatternFill

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color="1F4E78")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=9, color="595959")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=columns)
    ws["A3"] = note
    ws["A3"].font = Font(size=9, color="7F6000")
    ws["A3"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.row_dimensions[1].height = 26


def _table(
    ws: Any,
    columns: list[tuple[str, str, int, str]],
    rows: list[dict[str, Any]],
    *,
    header_row: int = 6,
    status_key: str | None = None,
) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    style = _styles()
    for index, (_key, header, width, _fmt) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=index, value=header)
        cell.fill = style["header_fill"]
        cell.font = style["header_font"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = style["border"]
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[header_row].height = 34
    status_colors = {
        "紧急补货": "F4CCCC",
        "关注": "FFF2CC",
        "滞销": "D9D9D9",
        "健康": "D9EAD3",
    }
    for row_index, row in enumerate(rows, start=header_row + 1):
        for col_index, (key, _header, _width, fmt) in enumerate(columns, start=1):
            value = row.get(key)
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.border = style["border"]
            cell.alignment = Alignment(vertical="top", wrap_text=key in {"name", "status", "action"})
            if fmt:
                cell.number_format = fmt
        if status_key:
            cell = ws.cell(row=row_index, column=next(i for i, c in enumerate(columns, start=1) if c[0] == status_key))
            fill = status_colors.get(str(row.get(status_key)))
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(bold=True)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(rows)}"
    )


def _write_metric_sheet(ws: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.append(["指标", "字段", "公式", "说明"])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    rows = [
        ("广告 CTR", "ctr", "点击 ÷ 展示", "广告点击率"),
        ("广告 CVR", "cvr", "广告订单 ÷ 点击", "广告转化率"),
        ("CPC", "cpc", "广告花费 ÷ 点击", "单次点击成本"),
        ("ACOS", "acos", "广告花费 ÷ 广告销售额", "广告成本占广告销售比"),
        ("ROAS", "roas", "广告销售额 ÷ 广告花费", "广告投入产出比"),
        ("广告贡献毛利", "ad_profit", "广告销售额 × 毛利率 - 广告花费", "广告口径的可变利润"),
        ("毛利率", "gross_margin", "毛利 ÷ 收入", "商品口径毛利，未扣平台佣金/履约/广告"),
        ("库存周转天数", "inventory_turnover_days", "平均库存件数 ÷ 日均销量", "30 天口径；越低周转越快"),
        ("库存周转次数", "inventory_turnover_times", "365 ÷ 库存周转天数", "年化周转次数"),
        ("库存覆盖天数", "days_cover", "(在库 + 在途) ÷ 日均销量", "含在途库存"),
        ("再订货点", "reorder_point", "日均销量 × (交期 + 安全天数)", "触发补货的库存水平"),
        ("建议补货量", "suggested_order_qty", "目标库存 - (在库 + 在途)", "目标覆盖 = 交期 + 安全 + 30 天"),
        ("ROI", "roi", "净利润 ÷ 总成本", "总成本=COGS+佣金+履约+广告+退款+仓储+其他"),
        ("净利率", "net_margin", "净利润 ÷ 收入", "含广告/退款/仓储等经营费用"),
    ]
    for row in rows:
        ws.append(row)
    widths = [16, 26, 48, 44]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def _save_with_properties(wb: Any, path: Path, title: str) -> None:
    wb.properties.title = title
    wb.properties.creator = "Outdoor ESS AI report_generator.py"
    wb.properties.description = "固定种子模拟经营数据；不接入真实后台。"
    utils.ensure_dir(path.parent)
    wb.save(path)
    logger.info("Excel 已生成: %s", path)


def write_ad_workbook(path: Path, ad_rows: list[dict[str, Any]], summary: dict[str, Any], run_id: str, period: str) -> None:
    _require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule

    wb = Workbook()
    ws = wb.active
    ws.title = "广告日报"
    columns = [
        ("date", "日期", 11, "yyyy-mm-dd"),
        ("weekday", "星期", 8, "@"),
        ("platform", "平台", 11, "@"),
        ("sku", "SKU", 20, "@"),
        ("name", "商品", 34, "@"),
        ("market", "市场", 8, "@"),
        ("impressions", "展示", 11, "#,##0"),
        ("clicks", "点击", 9, "#,##0"),
        ("ctr", "CTR", 8, "0.00%"),
        ("cpc", "CPC", 8, "$0.00"),
        ("spend", "广告花费", 12, "$#,##0.00"),
        ("ad_orders", "广告订单", 10, "#,##0"),
        ("cvr", "CVR", 8, "0.00%"),
        ("ad_sales", "广告销售额", 13, "$#,##0.00"),
        ("acos", "ACOS", 9, "0.00%"),
        ("roas", "ROAS", 8, "0.00"),
        ("ad_profit", "广告贡献毛利", 13, "$#,##0.00"),
        ("gross_margin_pct", "参考毛利率", 9, "0.0"),
    ]
    _meta(
        ws,
        "广告投放日报",
        f"运行：{run_id} ｜ 周期：{period} ｜ 记录：{len(ad_rows)} 条",
        "模拟数据：展示/点击/花费/订单均由固定种子生成；不接入广告后台。",
        len(columns),
    )
    _table(ws, columns, ad_rows)
    # 合计行
    total_row = 6 + len(ad_rows) + 1
    total = {
        "date": "合计",
        "platform": "",
        "sku": "",
        "name": "",
        "market": "",
        "impressions": sum(int(row["impressions"]) for row in ad_rows),
        "clicks": sum(int(row["clicks"]) for row in ad_rows),
        "spend": round(sum(float(row["spend"]) for row in ad_rows), 2),
        "ad_orders": sum(int(row["ad_orders"]) for row in ad_rows),
        "ad_sales": round(sum(float(row["ad_sales"]) for row in ad_rows), 2),
        "ad_profit": round(sum(float(row["ad_profit"]) for row in ad_rows), 2),
    }
    total["ctr"] = total["clicks"] / total["impressions"] if total["impressions"] else 0
    total["cvr"] = total["ad_orders"] / total["clicks"] if total["clicks"] else 0
    total["cpc"] = total["spend"] / total["clicks"] if total["clicks"] else 0
    total["acos"] = total["spend"] / total["ad_sales"] if total["ad_sales"] else 0
    total["roas"] = total["ad_sales"] / total["spend"] if total["spend"] else 0
    from openpyxl.styles import Font, PatternFill
    for index, (key, _header, _width, fmt) in enumerate(columns, start=1):
        cell = ws.cell(row=total_row, column=index, value=total.get(key, ""))
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
        if fmt:
            cell.number_format = fmt
    ws.auto_filter.ref = f"A6:R{6 + len(ad_rows)}"
    ws.conditional_formatting.add(
        f"O7:O{6 + len(ad_rows)}",
        ColorScaleRule(
            start_type="num", start_value=0, start_color="63BE7B",
            mid_type="num", mid_value=0.2, mid_color="FFEB84",
            end_type="num", end_value=0.5, end_color="F8696B",
        ),
    )
    metric_ws = wb.create_sheet("指标说明")
    _write_metric_sheet(metric_ws)
    _save_with_properties(wb, path, "广告投放日报")


def write_inventory_workbook(path: Path, metrics: list[dict[str, Any]], run_id: str, period: str) -> None:
    _require_openpyxl()
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "库存预警"
    rows = sorted(metrics, key=lambda item: item["days_cover"])
    columns = [
        ("sku", "SKU", 20, "@"),
        ("name", "商品", 32, "@"),
        ("platform", "平台", 11, "@"),
        ("market", "市场", 8, "@"),
        ("units", "30天销量", 10, "#,##0"),
        ("avg_daily_units", "日均销量", 10, "0.00"),
        ("begin_inventory", "期初库存", 10, "#,##0"),
        ("on_hand", "在库", 9, "#,##0"),
        ("in_transit", "在途", 9, "#,##0"),
        ("days_cover", "覆盖天数", 10, "0.0"),
        ("lead_time_days", "交期(天)", 9, "0"),
        ("safety_days", "安全天数", 9, "0"),
        ("reorder_point", "再订货点", 11, "0.0"),
        ("suggested_order_qty", "建议补货", 11, "#,##0"),
        ("stock_value", "库存金额", 13, "$#,##0.00"),
        ("aging_units", "库龄>90天", 11, "#,##0"),
        ("inventory_turnover_days", "周转天数", 10, "0.0"),
        ("inventory_turnover_times", "年化周转", 10, "0.00"),
        ("stock_status", "状态", 10, "@"),
        ("priority", "优先级", 9, "@"),
    ]
    _meta(
        ws,
        "库存预警表",
        f"运行：{run_id} ｜ 周期：{period} ｜ SKU：{len(metrics)}",
        "覆盖天数=(在库+在途)÷日均销量；周转天数=平均库存÷日均销量；模拟数据。",
        len(columns),
    )
    _table(ws, columns, rows, status_key="stock_status")
    ws.auto_filter.ref = f"A6:T{6 + len(rows)}"
    metric_ws = wb.create_sheet("指标说明")
    _write_metric_sheet(metric_ws)
    _save_with_properties(wb, path, "库存预警表")


def write_profit_workbook(path: Path, metrics: list[dict[str, Any]], summary: dict[str, Any], run_id: str, period: str) -> None:
    _require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "经营概览"
    _meta(
        dashboard,
        "月度经营概览",
        f"运行：{run_id} ｜ 周期：{period} ｜ SKU：{summary['sku_count']}",
        "模拟数据：收入/成本/广告/库存均由固定种子生成；不接入真实后台。",
        8,
    )
    dashboard["A5"] = "核心指标"
    dashboard["A5"].font = Font(bold=True, size=12, color="1F4E78")
    overview = [
        ("总销量", summary["units"], "#,##0"),
        ("总收入", summary["revenue"], "$#,##0.00"),
        ("总广告花费", summary["ad_spend"], "$#,##0.00"),
        ("总净利润", summary["net_profit"], "$#,##0.00"),
        ("净利率", summary["net_margin"] / 100, "0.00%"),
        ("整体 ROI", summary["roi"] / 100, "0.00%"),
        ("广告 ACOS", summary["ad_acos"] / 100, "0.00%"),
        ("广告 ROAS", summary["ad_roas"], "0.00"),
        ("库存预警 SKU", summary["inventory_alerts"], "#,##0"),
        ("紧急补货 SKU", summary["urgent_alerts"], "#,##0"),
    ]
    for index, (label, value, fmt) in enumerate(overview, start=6):
        dashboard.cell(row=index, column=1, value=label).font = Font(bold=True)
        cell = dashboard.cell(row=index, column=2, value=value)
        cell.number_format = fmt
        cell.alignment = Alignment(horizontal="right")
    dashboard["D5"] = "Top / 风险 SKU"
    dashboard["D5"].font = Font(bold=True, size=12, color="1F4E78")
    top_roi = next((item for item in metrics if item["sku"] == summary["top_roi"]), None)
    low_cover = next((item for item in metrics if item["sku"] == summary["lowest_cover"]), None)
    dashboard["D6"] = "最高 ROI"
    dashboard["E6"] = f"{top_roi['sku'] if top_roi else '-'} / {top_roi['roi'] * 100:.1f}%" if top_roi else "-"
    dashboard["D7"] = "最低覆盖"
    dashboard["E7"] = f"{low_cover['sku'] if low_cover else '-'} / {low_cover['days_cover']:.1f} 天" if low_cover else "-"
    dashboard["D8"] = "建议动作"
    dashboard["E8"] = "优先处理 P1 补货，复核滞销库存与广告 ACOS"
    dashboard.column_dimensions["A"].width = 22
    dashboard.column_dimensions["B"].width = 18
    dashboard.column_dimensions["D"].width = 22
    dashboard.column_dimensions["E"].width = 46

    ws = wb.create_sheet("月度利润")
    columns = [
        ("sku", "SKU", 20, "@"),
        ("name", "商品", 32, "@"),
        ("platform", "平台", 11, "@"),
        ("market", "市场", 8, "@"),
        ("units", "销量", 9, "#,##0"),
        ("revenue", "收入", 13, "$#,##0.00"),
        ("cogs", "COGS", 13, "$#,##0.00"),
        ("gross_profit", "毛利", 13, "$#,##0.00"),
        ("gross_margin", "毛利率(商品)", 11, "0.00%"),
        ("platform_fee", "平台佣金", 12, "$#,##0.00"),
        ("fulfillment", "履约", 11, "$#,##0.00"),
        ("ad_spend", "广告费", 11, "$#,##0.00"),
        ("refunds", "退款", 11, "$#,##0.00"),
        ("storage", "仓储", 10, "$#,##0.00"),
        ("other", "其他", 10, "$#,##0.00"),
        ("total_cost", "总成本", 13, "$#,##0.00"),
        ("net_profit", "净利润", 13, "$#,##0.00"),
        ("net_margin", "净利率", 10, "0.00%"),
        ("roi", "ROI", 9, "0.00%"),
        ("ad_acos", "广告 ACOS", 11, "0.00%"),
        ("ad_roas", "广告 ROAS", 10, "0.00"),
        ("avg_inventory_units", "平均库存", 11, "#,##0.00"),
        ("inventory_turnover_days", "库存周转天数", 12, "0.0"),
        ("inventory_turnover_times", "年化周转", 10, "0.00"),
        ("stock_value", "库存金额", 12, "$#,##0.00"),
    ]
    _meta(
        ws,
        "月度利润核算表",
        f"运行：{run_id} ｜ 周期：{period} ｜ 金额单位：USD",
        "ROI=净利润÷总成本；总成本=COGS+佣金+履约+广告+退款+仓储+其他；库存周转天数=平均库存÷日均销量。",
        len(columns),
    )
    _table(ws, columns, metrics)
    total_row = 6 + len(metrics) + 1
    totals = {
        "sku": "合计",
        "name": "",
        "platform": "",
        "market": "",
        "units": summary["units"],
        "revenue": summary["revenue"],
        "cogs": summary["cogs"],
        "gross_profit": round(summary["revenue"] - summary["cogs"], 2),
        "platform_fee": round(sum(item["platform_fee"] for item in metrics), 2),
        "fulfillment": round(sum(item["fulfillment"] for item in metrics), 2),
        "ad_spend": summary["ad_spend"],
        "refunds": round(sum(item["refunds"] for item in metrics), 2),
        "storage": round(sum(item["storage"] for item in metrics), 2),
        "other": round(sum(item["other"] for item in metrics), 2),
        "total_cost": round(sum(item["total_cost"] for item in metrics), 2),
        "net_profit": summary["net_profit"],
        "avg_inventory_units": round(sum(item["avg_inventory_units"] for item in metrics), 2),
        "stock_value": round(sum(item["stock_value"] for item in metrics), 2),
    }
    totals["gross_margin"] = totals["gross_profit"] / totals["revenue"] if totals["revenue"] else 0
    totals["net_margin"] = totals["net_profit"] / totals["revenue"] if totals["revenue"] else 0
    totals["roi"] = totals["net_profit"] / sum(item["total_cost"] for item in metrics) if metrics else 0
    totals["ad_acos"] = summary["ad_acos"] / 100
    totals["ad_roas"] = summary["ad_roas"]
    avg_daily_total = sum(item["avg_daily_units"] for item in metrics)
    totals["inventory_turnover_days"] = (
        round(sum(item["avg_inventory_units"] for item in metrics) / avg_daily_total, 1)
        if avg_daily_total
        else 0
    )
    totals["inventory_turnover_times"] = (
        round(365 / totals["inventory_turnover_days"], 2) if totals["inventory_turnover_days"] else 0
    )
    for index, (key, _header, _width, fmt) in enumerate(columns, start=1):
        cell = ws.cell(row=total_row, column=index, value=totals.get(key, ""))
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
        if fmt:
            cell.number_format = fmt
    ws.auto_filter.ref = f"A6:Y{6 + len(metrics)}"
    ws.conditional_formatting.add(
        f"S7:S{6 + len(metrics)}",
        ColorScaleRule(
            start_type="num", start_value=-0.2, start_color="F8696B",
            mid_type="num", mid_value=0.2, mid_color="FFEB84",
            end_type="num", end_value=0.6, end_color="63BE7B",
        ),
    )
    metric_ws = wb.create_sheet("指标说明")
    _write_metric_sheet(metric_ws)
    _save_with_properties(wb, path, "月度利润核算表")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def _display(root: Path, path: Path) -> str:
    try:
        return str(path.resolve().relative_to(root.resolve()))
    except ValueError:
        return str(path.resolve())


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="report_generator",
        description="模拟经营数据：广告日报、库存预警、月度利润报表",
    )
    parser.add_argument("--config", default=str(DEFAULT_CONFIG), help="报表配置")
    parser.add_argument("--input", default=str(DEFAULT_INPUT), help="选品表")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT), help="输出目录")
    parser.add_argument("--seed", type=int, help="覆盖模拟种子")
    parser.add_argument("--period-end", help="报表截止日 YYYY-MM-DD")
    parser.add_argument("--days", type=int, help="报表周期天数")
    parser.add_argument("--log-level", default="INFO")
    return parser


def _configure_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if callable(reconfigure):
            reconfigure(encoding="utf-8", errors="replace")


def main(argv: Iterable[str] | None = None) -> int:
    _configure_stdio()
    args = build_parser().parse_args(list(argv) if argv is not None else None)
    try:
        config = load_config(args.config)
        if args.seed is not None:
            config.data.setdefault("simulation", {})["seed"] = args.seed
        if args.period_end:
            config.data.setdefault("simulation", {})["period_end"] = args.period_end
        if args.days is not None:
            config.data.setdefault("simulation", {})["days"] = args.days
        products = load_products(args.input, str(config.get("input.sheet", "高潜力选品")))
        output_dir = Path(args.output_dir).expanduser().resolve()
        utils.ensure_dir(output_dir)
        utils.setup_logging(
            level=args.log_level,
            log_file=output_dir / "report_generator.log",
            console=True,
            force=True,
        )
        simulator = OperationsSimulator(config, products)
        data = simulator.run()
        run_id = f"ops_{utils.utc_now().strftime('%Y%m%d_%H%M%S_%f')}"
        ad_path = output_dir / str(config.get("output.ad_file", "ad_daily_report.xlsx"))
        inventory_path = output_dir / str(config.get("output.inventory_file", "inventory_alert.xlsx"))
        profit_path = output_dir / str(config.get("output.profit_file", "monthly_profit.xlsx"))
        write_ad_workbook(ad_path, data["ad_rows"], data["summary"], run_id, data["period"])
        write_inventory_workbook(inventory_path, data["metrics"], run_id, data["period"])
        write_profit_workbook(profit_path, data["metrics"], data["summary"], run_id, data["period"])
        summary = {
            **data["summary"],
            "run_id": run_id,
            "generated_at": utils.utc_now_iso(),
            "period": data["period"],
            "days": data["days"],
            "seed": data["seed"],
            "notice": "固定种子模拟经营数据；不接入真实店铺/广告/库存后台。",
            "metrics": data["metrics"],
            "outputs": {
                "ad": _display(config.root, ad_path),
                "inventory": _display(config.root, inventory_path),
                "profit": _display(config.root, profit_path),
                "summary": _display(config.root, output_dir / str(config.get("output.summary_file", "report_summary.json"))),
            },
        }
        utils.write_json(output_dir / str(config.get("output.summary_file", "report_summary.json")), summary)
    except (ReportError, utils.UtilError, OSError, ValueError) as exc:
        logger.exception("报表生成失败: %s", exc)
        print(f"✗ 报表生成失败：{exc}", file=sys.stderr)
        return 1

    print("")
    print("=" * 72)
    print(f"✓ 经营报表生成完成：{run_id}")
    print(f"  周期    : {data['period']}（{data['days']} 天）")
    print(f"  收入    : ${data['summary']['revenue']:,.2f}")
    print(
        f"  净利润  : ${data['summary']['net_profit']:,.2f}  |  "
        f"ROI {data['summary']['roi']:.2f}%  |  ACOS {data['summary']['ad_acos']:.2f}%"
    )
    print(f"  库存预警: {data['summary']['inventory_alerts']} 个 SKU")
    print(f"  广告日报: {summary['outputs']['ad']}")
    print(f"  库存预警: {summary['outputs']['inventory']}")
    print(f"  月度利润: {summary['outputs']['profit']}")
    print("=" * 72)
    return 0


if __name__ == "__main__":
    from error_guard import run_cli

    raise SystemExit(run_cli(main, "report_generator"))
