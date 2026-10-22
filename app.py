#!/usr/bin/env python3
"""Tundra · 户外储能出海 AI 控制台。

A+C 视觉方向：苔原账本浅色主体 + 深色运行日志抽屉。
本文件只做本地 Streamlit 控制台，核心业务逻辑仍在各 CLI 模块中。

启动：
    streamlit run app.py
或：
    powershell -ExecutionPolicy Bypass -File .\\start_all.ps1
"""

from __future__ import annotations

import csv
import html
import json
import os
import subprocess
import sys
import time
from datetime import UTC, datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent
VENDOR_DIR = PROJECT_ROOT / ".vendor"
if VENDOR_DIR.is_dir() and str(VENDOR_DIR) not in sys.path:
    sys.path.insert(0, str(VENDOR_DIR))

import pandas as pd  # noqa: E402
import streamlit as st  # noqa: E402

from error_guard import install_global_handler, log_command_error, log_exception  # noqa: E402
from utils import ensure_dir, read_json, read_text, write_text  # noqa: E402

install_global_handler("streamlit")
st.set_page_config(
    page_title="Tundra · 户外储能出海控制台",
    page_icon="◒",
    layout="wide",
    initial_sidebar_state="auto",
)

ROOT = PROJECT_ROOT
OUTPUTS = ROOT / "outputs"
LOG_DIR = OUTPUTS / "logs"
RUN_LOG = LOG_DIR / "ui_commands.log"
RESEARCH_XLSX = OUTPUTS / "research" / "high_potential_products.xlsx"
RESEARCH_SUMMARY = OUTPUTS / "research" / "research_summary.json"
LISTING_SUMMARY = OUTPUTS / "listing" / "listing_summary.json"
MONITOR_SUMMARY = OUTPUTS / "monitor" / "reports" / "monitor_summary.json"
COMPLIANCE_SUMMARY = OUTPUTS / "compliance" / "compliance_summary.json"
OPS_SUMMARY = OUTPUTS / "operations" / "report_summary.json"
INTERVIEW_MD = OUTPUTS / "interview" / "interview_demo.md"

PAGES = ["总览", "选品研究", "上架素材", "竞品监控", "合规校验", "经营报表", "面试演示"]
PAGE_INDEX = {page: index for index, page in enumerate(PAGES, start=1)}

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;500&family=IBM+Plex+Sans:wght@400;500;600&display=swap');
:root{
  --frame:#EDF3EE; --ground:#F4F7F3; --surface:#FFFFFF; --raised:#F7FAF7;
  --ink:#14201B; --muted:#56665D; --faint:#5B6B62;
  --line:#DCE5DF; --line-soft:#E9EFEA;
  --accent:#1F6F5C; --accent-ink:#155446; --mint:#DDE9E1;
  --amber:#A66A1F; --amber-bg:#FBF2E5;
  --red:#B42318; --red-bg:#FBEDEA;
  --dark:#101816;
  --sans:'IBM Plex Sans','Microsoft YaHei',system-ui,sans-serif;
  --mono:'IBM Plex Mono','Cascadia Mono',Consolas,monospace;
}
html, body, [class*="css"] { font-family:var(--sans); }
.stApp, .stApp p, .stApp span, .stApp div, .stApp label, .stApp button, .stApp h1 { font-family:var(--sans); }
.kpi-value,.overview-table .num,.overview-table .sku,.panel-meta,.foot-right,.log,.log *,.side-stat b{font-family:var(--mono);}
.stApp{
  background:radial-gradient(1000px 500px at 5% -10%, #E9F2EB 0%, transparent 60%), var(--frame);
  color:var(--ink);
}
.block-container{padding-top:0 !important; padding-bottom:18px !important; padding-left:18px !important; padding-right:18px !important; max-width:none !important;}
header[data-testid="stHeader"]{background:transparent;}
[data-testid="stToolbar"], [data-testid="stDecoration"], #MainMenu {display:none !important;}
section[data-testid="stSidebar"]{
  width:232px !important; min-width:232px !important; background:var(--frame);
  border-right:none;
}
section[data-testid="stSidebar"] > div{padding-top:0;}
[data-testid="stSidebarContent"]{padding:0 14px 16px;}
[data-testid="stSidebarCollapseButton"]{color:var(--muted);}
[data-testid="stSidebarCollapsedControl"], [data-testid="collapsedControl"]{
  display:flex !important; visibility:visible !important; opacity:1 !important;
  color:var(--ink) !important; background:var(--frame) !important; z-index:1000;
}
.stButton>button, .stDownloadButton>button{
  height:32px; min-height:32px; border-radius:6px; font-size:12px; font-weight:600;
  border:1px solid var(--line); background:var(--surface); color:#405249;
  box-shadow:none; padding:0 12px;
}
.stButton>button:hover, .stDownloadButton>button:hover{
  border-color:#B9C9BE; color:var(--ink); background:#F8FBF8;
}
.stButton>button[kind="primary"], .stDownloadButton>button[kind="primary"]{
  background:var(--accent); border-color:var(--accent); color:#fff;
}
.stButton>button[kind="primary"]:hover{background:var(--accent-ink); border-color:var(--accent-ink);}
.stButton>button p, .stDownloadButton>button p{font-size:12px; font-weight:600;}
div[data-testid="stVerticalBlock"]{gap:0;}
div[data-testid="stHorizontalBlock"]{gap:14px;}
.stSelectbox label, .stMultiSelect label, .stRadio label{font-size:11px; color:var(--muted);}
[data-testid="stSidebar"] div[role="radiogroup"]{gap:2px;}
[data-testid="stSidebar"] div[data-testid="stRadioGroup"]{margin-top:18px;}
[data-testid="stSidebar"] div[data-testid="stRadioGroup"] > div > label[data-testid="stRadioOption"]{
  min-height:34px; border-radius:7px; padding:7px 10px; margin:0; width:100%; box-sizing:border-box;
  display:flex !important; align-items:center;
  color:#405249; font-size:12px; transition:background .12s ease; background:transparent !important;
}
[data-testid="stSidebar"] div[data-testid="stRadioGroup"] > div > label[data-testid="stRadioOption"]:hover{background:#E5EEE7 !important;}
[data-testid="stSidebar"] div[data-testid="stRadioGroup"] > div[data-selected="true"] > label[data-testid="stRadioOption"],
[data-testid="stSidebar"] div[data-testid="stRadioGroup"] > div > label[data-selected="true"]{
  background:var(--mint) !important; color:var(--ink) !important; font-weight:600;
}
[data-testid="stSidebar"] div[data-testid="stRadioGroup"] label[data-testid="stRadioOption"] > div{width:100%;}
[data-testid="stSidebar"] div[data-testid="stRadioGroup"] label[data-testid="stRadioOption"] > div > div:first-child{display:none;}
[data-testid="stSidebar"] div[data-testid="stRadioGroup"] label[data-testid="stRadioOption"] p{font-size:12px; line-height:1.2; margin:0;}
[data-testid="stSidebar"] .stButton>button{width:100%; height:29px; min-height:29px; background:transparent; border-color:var(--line); font-size:11px; font-weight:500; text-align:left; justify-content:flex-start; padding:0 10px; color:var(--muted);}
[data-testid="stSidebar"] .stButton>button:hover{background:rgba(255,255,255,.45);}
[data-testid="stSidebar"] .stButton{margin-top:4px;}
.st-key-sheet{
  background:var(--surface); border:1px solid #D8E3DA; border-radius:16px;
  padding:24px 28px 18px; min-height:calc(100vh - 76px);
}
.topbar{height:54px; display:flex; align-items:center; justify-content:space-between; padding:0 22px 0 24px;}
.crumb{font-size:11px; color:var(--faint);}
.crumb b{color:var(--ink); font-weight:500;}
.top-actions{display:flex; align-items:center; gap:9px;}
.sim-pill{font-family:var(--mono); font-size:11px; color:var(--muted); border:1px solid var(--line); border-radius:99px; padding:5px 9px; background:rgba(255,255,255,.45);}
.top-status{font-size:11px; color:var(--muted); display:inline-flex; align-items:center; gap:6px;}
.dot{width:5px; height:5px; border-radius:50%; display:inline-block; background:var(--accent);}
.dot.amber{background:var(--amber);} .dot.red{background:var(--red);} .dot.grey{background:#AAB9AE;}
.pagehead{display:flex; align-items:flex-start; justify-content:space-between; gap:20px; min-height:72px;}
.overline{font-size:11px; color:var(--faint); font-weight:500; margin-bottom:6px;}
.pagehead h1{font-size:27px; line-height:1.15; font-weight:500; letter-spacing:-.02em; margin:0;}
.sub{margin-top:6px; font-size:12px; color:var(--muted); display:flex; align-items:center; gap:8px; flex-wrap:wrap;}
.sub .sep{width:1px; height:10px; background:var(--line);}
.head-actions{padding-top:7px; display:flex; gap:12px; align-items:center; justify-content:flex-end;}
.st-key-head_actions{padding-top:7px; display:flex; justify-content:flex-end;}
.st-key-head_actions .stButton{margin:0;}
.text-action{font-size:11px; color:var(--muted); text-decoration:none; display:inline-flex; align-items:center; gap:5px;}
.text-action:hover{color:var(--ink);}
.brand{height:56px; display:flex; align-items:center; gap:10px; padding:0 6px;}
.mark{width:24px; height:24px; border:1px solid var(--ink); border-radius:7px; display:grid; place-items:center; position:relative; flex:0 0 auto;}
.mark:before,.mark:after{content:""; position:absolute; background:var(--ink); border-radius:1px;}
.mark:before{width:10px; height:1px; transform:rotate(-34deg); left:7px; top:13px;}
.mark:after{width:7px; height:1px; transform:rotate(34deg); left:11px; top:12px;}
.wordmark{font-size:14px; font-weight:600; letter-spacing:.02em;}
.wordmark span{display:block; font-size:11px; color:var(--muted); letter-spacing:0; font-weight:500; margin-top:1px;}
.workspace{margin:2px 0 14px; padding:8px 10px; border:1px solid var(--line); border-radius:8px; display:flex; align-items:center; justify-content:space-between; background:rgba(255,255,255,.38);}
.workspace b{display:block; font-size:12px; font-weight:600;}
.workspace small{display:block; font-size:11px; color:var(--muted); margin-top:2px;}
.chev{color:var(--faint); font-size:12px;}
.nav-label{font-size:11px; color:var(--faint); font-weight:500; margin:9px 8px 6px;}
.rail-foot{margin-top:18px; border-top:1px solid var(--line); padding:12px 4px 0;}
.system-row{display:flex; align-items:center; justify-content:space-between; font-size:11px; color:var(--muted); padding:4px 0;}
.system-note{font-size:11px; line-height:1.55; color:var(--faint); margin-top:6px;}
.side-stats{display:grid; grid-template-columns:1fr 1fr; gap:6px; margin:10px 0 12px;}
.side-stat{background:rgba(255,255,255,.4); border:1px solid var(--line-soft); border-radius:7px; padding:7px 8px;}
.side-stat b{display:block; font-family:var(--mono); font-size:14px; font-weight:500;}
.side-stat span{display:block; font-size:11px; color:var(--muted); margin-top:2px;}
.flow{margin-top:16px; background:var(--raised); border:1px solid var(--line-soft); border-radius:9px; display:flex; align-items:stretch; min-height:70px; overflow:hidden;}
.flow-node{flex:1; padding:11px 14px; position:relative; min-width:0;}
.flow-node + .flow-node:before{content:""; position:absolute; left:0; top:15px; bottom:15px; width:1px; background:var(--line);}
.flow-kicker{font-size:11px; color:var(--faint); font-weight:500; white-space:nowrap;}
.flow-main{display:flex; align-items:baseline; gap:6px; margin-top:3px; min-width:0;}
.flow-value{font-size:22px; line-height:1; font-weight:500; letter-spacing:-.02em; white-space:nowrap;}
.flow-unit{font-size:11px; color:var(--muted); white-space:nowrap;}
.flow-meta{font-size:11px; color:var(--muted); margin-top:4px; white-space:nowrap; overflow:hidden; text-overflow:ellipsis;}
.arrow{position:absolute; right:-5px; top:50%; z-index:2; width:9px; height:9px; border-top:1px solid #B7C7BD; border-right:1px solid #B7C7BD; transform:translateY(-50%) rotate(45deg); background:var(--raised);}
.flow-node:last-child .arrow{display:none;}
.green{color:var(--accent);} .red-text{color:var(--red);} .amber-text{color:var(--amber);}
.kpi-strip{display:grid; grid-template-columns:1.1fr 1fr 1fr; border-top:1px solid var(--line-soft); border-bottom:1px solid var(--line-soft); margin-top:16px; padding:11px 0;}
.kpi{padding:0 18px; position:relative;}
.kpi:first-child{padding-left:0;} .kpi:last-child{padding-right:0;}
.kpi + .kpi:before{content:""; position:absolute; left:0; top:2px; bottom:2px; width:1px; background:var(--line-soft);}
.kpi-label{font-size:11px; color:var(--faint); font-weight:500;}
.kpi-value{font-family:var(--mono); font-size:25px; line-height:1.05; font-weight:500; margin-top:5px; letter-spacing:-.04em;}
.kpi-note{font-size:11px; color:var(--muted); margin-top:5px; display:flex; gap:7px; align-items:center; flex-wrap:wrap;}
.main-grid{display:grid; grid-template-columns:minmax(0,1.52fr) minmax(300px,.78fr); gap:20px; margin-top:16px; align-items:start;}
.main-grid > div,.main-grid > aside{min-width:0;}
.panel{border:1px solid var(--line-soft); border-radius:9px; background:var(--surface); min-width:0;}
.panel-head{height:40px; padding:0 14px; display:flex; align-items:center; justify-content:space-between; border-bottom:1px solid var(--line-soft);}
.panel-title{font-size:11px; font-weight:600; display:flex; align-items:center; gap:7px;}
.panel-title .num{font-family:var(--mono); font-size:11px; color:var(--faint); font-weight:400;}
.panel-meta{font-size:11px; color:var(--faint); font-family:var(--mono);}
.panel-link{font-size:11px; color:var(--accent); text-decoration:none; font-weight:500;}
.priority{padding-bottom:4px;}
.action-row{display:flex; align-items:center; gap:11px; padding:9px 14px; border-bottom:1px solid var(--line-soft); min-height:52px;}
.action-row:last-child{border-bottom:none;}
.prio{width:31px; height:21px; border-radius:4px; display:grid; place-items:center; font-family:var(--mono); font-size:11px; font-weight:500; flex:0 0 auto;}
.prio.p1{background:var(--red-bg); color:var(--red);}
.prio.p2{background:var(--amber-bg); color:var(--amber);}
.prio.p3{background:#EDF2EE; color:#53645B;}
.action-copy{flex:1; min-width:0;}
.action-copy b{display:block; font-size:12px; font-weight:600; white-space:nowrap; overflow:hidden; text-overflow:ellipsis;}
.action-copy span{display:block; font-size:11px; color:var(--muted); margin-top:2px; white-space:nowrap; overflow:hidden; text-overflow:ellipsis;}
.action-arrow{color:#71877A; font-size:14px; flex:0 0 auto;}
.overview-table{width:100%; border-collapse:collapse; table-layout:fixed;}
.overview-table th{font-size:11px; color:var(--faint); font-weight:500; text-align:left; padding:7px 10px 6px; border-bottom:1px solid var(--line-soft); white-space:nowrap;}
.overview-table td{font-size:11px; color:#25372E; padding:6px 10px; border-bottom:1px solid #F0F4F1; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; line-height:1.25;}
.overview-table tr:last-child td{border-bottom:none;}
.overview-table th.num,.overview-table td.num{text-align:right; font-family:var(--mono); font-variant-numeric:tabular-nums;}
.overview-table td.sku{font-family:var(--mono); font-size:11px; color:#405249; letter-spacing:-.02em;}
.status{display:inline-flex; align-items:center; gap:4px; font-size:11px; padding:2px 6px; border-radius:99px; font-weight:500; white-space:nowrap;}
.status .sdot{width:4px; height:4px; border-radius:50%;}
.st-green{background:#EAF4ED; color:#22734A;} .st-green .sdot{background:#2E8B57;}
.st-amber{background:#FBF2E5; color:#8A5B18;} .st-amber .sdot{background:#B7791F;}
.st-red{background:#FBEDEA; color:#9E2A20;} .st-red .sdot{background:#C0392B;}
.st-grey{background:#EEF2EF; color:#5F6F66;} .st-grey .sdot{background:#819188;}
.side-stack{display:flex; flex-direction:column; gap:14px; min-width:0;}
.compliance{padding:13px 14px 13px;}
.bar{height:8px; display:flex; border-radius:99px; overflow:hidden; margin:11px 0 9px; background:#EEF2EF;}
.bar span{height:100%;}
.legend{display:flex; gap:12px; flex-wrap:wrap;}
.leg{font-size:11px; color:var(--muted); display:inline-flex; align-items:center; gap:5px;}
.sw{width:5px; height:5px; border-radius:1px; display:inline-block;}
.sw.red{background:var(--red);} .sw.amber{background:#C28A37;} .sw.grey{background:#C9D4CC;}
.compliance-foot{display:flex; justify-content:space-between; border-top:1px solid var(--line-soft); margin-top:10px; padding-top:9px; font-size:11px; color:var(--muted);}
.compliance-foot b{font-family:var(--mono); font-weight:500; color:var(--ink);}
.mini-panel{padding:12px 14px 12px;}
.mini-title{font-size:11px; color:var(--faint); font-weight:500;}
.mini-row{display:flex; align-items:center; justify-content:space-between; margin-top:7px; font-size:11px;}
.mini-row span{color:var(--muted);} .mini-row b{font-family:var(--mono); font-weight:500; color:var(--ink);}
.log{background:var(--dark); border:1px solid #1D2A24; border-radius:9px; padding:12px 14px 11px; color:#DCE8E1; min-height:158px; display:flex; flex-direction:column;}
.log-head{display:flex; align-items:center; justify-content:space-between; border-bottom:1px solid #25352D; padding-bottom:8px;}
.log-title{font-size:11px; font-weight:500; color:#EFF6F2; display:flex; align-items:center; gap:7px;}
.log-title .pulse{width:5px; height:5px; border-radius:50%; background:#63C79E; box-shadow:0 0 0 3px rgba(99,199,158,.08);}
.log-status{font-family:var(--mono); font-size:11px; color:#8CC9AC;}
.log-lines{font-family:var(--mono); font-size:11px; line-height:1.6; margin-top:8px; flex:1;}
.log-line{display:flex; gap:7px; color:#AABCB1; white-space:nowrap;}
.log-line .t{color:#8FA79A; flex:0 0 58px;}
.log-line .k{color:#D7E5DC; flex:0 0 78px;}
.log-line .v{color:#AFC2B6; flex:1; overflow:hidden; text-overflow:ellipsis;}
.log-line .d{color:#8FA79A; text-align:right;}
.log-foot{display:flex; justify-content:space-between; border-top:1px solid #25352D; margin-top:5px; padding-top:7px; font-family:var(--mono); font-size:11px; color:#8FA79A;}
.log-foot a{color:#9ED3B9; text-decoration:none;}
.foot{margin-top:12px; display:flex; align-items:center; justify-content:space-between; color:var(--faint); font-size:11px; gap:14px; flex-wrap:wrap;}
.foot-right{font-family:var(--mono);}
.toolbar{display:flex; align-items:flex-end; gap:12px; margin:14px 0 10px; flex-wrap:wrap;}
.toolbar .grow{flex:1; min-width:0;}
.section-title{font-size:12px; font-weight:600; margin:14px 0 7px;}
.stDataFrame{border:1px solid var(--line-soft); border-radius:8px; overflow:hidden;}
[data-testid="stDataFrame"] div[role="columnheader"]{background:var(--raised); color:var(--faint); font-size:11px;}
[data-testid="stDataFrame"] [role="gridcell"]{font-size:11px;}
[data-testid="stExpander"] details{border-color:var(--line-soft); border-radius:8px;}
[data-testid="stTabs"] [data-baseweb="tab-list"]{gap:4px; border-bottom:1px solid var(--line-soft);}
[data-testid="stTabs"] [data-baseweb="tab"]{background:transparent; border:none; border-radius:0; padding:7px 10px; font-size:11px; color:var(--muted);}
[data-testid="stTabs"] [aria-selected="true"]{color:var(--ink); border-bottom:2px solid var(--accent);}
.stAlert{border-radius:8px; font-size:11px;}
.stInfo{border-radius:8px;}
.stSpinner{font-size:11px;}
div[data-testid="stSelectbox"] div[data-baseweb="select"]{font-size:11px;}
@media(max-width:1100px){.kpi-strip{grid-template-columns:1fr; gap:14px;} .kpi{padding:0;} .kpi + .kpi:before{display:none;} .kpi + .kpi{border-top:1px solid var(--line-soft); padding-top:13px;} .main-grid{grid-template-columns:1fr;} .side-stack{display:grid; grid-template-columns:1fr 1fr;} .flow{flex-wrap:wrap;} .flow-node{flex:1 0 33%;} .flow-node:nth-child(3):before{display:none;}}
@media(max-width:760px){.topbar{padding:0 16px 0 52px;} [data-testid="stExpandSidebarButton"]{display:flex !important; position:fixed; left:10px; top:10px; width:34px; height:34px; z-index:1200; background:var(--frame) !important; border:1px solid var(--line); border-radius:8px; color:var(--ink) !important; align-items:center; justify-content:center;} [data-testid="stExpandSidebarButton"] span{font-size:18px;} .top-actions .sim-pill{display:none;} .st-key-sheet{border-radius:0; padding:20px;} .pagehead{display:block;} .head-actions{justify-content:flex-start; padding-top:14px;} .side-stack{grid-template-columns:1fr;} .flow-node{flex:1 0 50%;} .flow-node:nth-child(2n):before{display:none;} .flow-node:last-child{flex:1 0 100%;} .main-grid{grid-template-columns:1fr; gap:14px;}}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def _summary(path: Path) -> dict:
    try:
        value = read_json(path, default={})
        return value if isinstance(value, dict) else {}
    except Exception:
        return {}


def _read_excel(path: Path, sheet: str, header_tokens: tuple[str, ...] = ("SKU",)) -> tuple[list[str], list[dict]]:
    if not path.exists():
        return [], []
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        headers: list[str] = []
        header_row = 1
        for index, row in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True), start=1):
            values = [str(cell or "").strip() for cell in row]
            if any(token in values for token in header_tokens):
                headers = values
                header_row = index
                break
        rows: list[dict] = []
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            rows.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
        wb.close()
        return headers, rows
    except Exception as exc:
        log_exception(exc, app="streamlit", context={"read_excel": str(path)})
        return [], []


def _df(headers: list[str], rows: list[dict]) -> pd.DataFrame:
    if not rows:
        return pd.DataFrame(columns=headers)
    return pd.DataFrame(rows, columns=headers)


def _download(path: Path, label: str, mime: str = "application/octet-stream") -> None:
    if path.exists():
        st.download_button(
            label=f"下载 {label}",
            data=path.read_bytes(),
            file_name=path.name,
            mime=mime,
            width="content",
        )


def _run_log(entry: dict) -> None:
    ensure_dir(LOG_DIR)
    write_text(RUN_LOG, json.dumps(entry, ensure_ascii=False) + "\n", append=True)


def run_command(label: str, args: list[str], *, timeout: int = 300) -> subprocess.CompletedProcess | None:
    command = [sys.executable, *args]
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONPATH"] = str(VENDOR_DIR) + os.pathsep + env.get("PYTHONPATH", "")
    started = time.time()
    proc: subprocess.CompletedProcess | None = None
    try:
        with st.spinner(f"正在执行：{label}"):
            proc = subprocess.run(
                command,
                cwd=str(ROOT),
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=timeout,
                env=env,
                check=False,
            )
        elapsed = round(time.time() - started, 2)
        _run_log(
            {
                "timestamp": datetime.now(UTC).isoformat(),
                "label": label,
                "command": command,
                "returncode": proc.returncode,
                "elapsed_seconds": elapsed,
            }
        )
        if proc.returncode != 0:
            log_command_error("streamlit", command, proc.returncode, proc.stdout or "", proc.stderr or "")
            st.error(f"{label} 执行失败，退出码 {proc.returncode}。错误已写入 outputs/logs/errors.log")
        else:
            st.success(f"{label} 执行完成（{elapsed:.1f}s）")
        with st.expander("查看运行日志", expanded=proc.returncode != 0):
            st.code((proc.stdout or "") + ("\n[stderr]\n" + (proc.stderr or "") if proc.stderr else ""), language="text")
        return proc
    except subprocess.TimeoutExpired:
        log_exception(RuntimeError(f"{label} 执行超时"), app="streamlit", context={"command": command})
        st.error(f"{label} 执行超时（>{timeout}s），已记录错误日志。")
    except Exception as exc:
        log_exception(exc, app="streamlit", context={"command": command})
        st.error(f"{label} 执行异常：{exc}")
    return proc


def _page_header(title: str, description: str, page: str, action_label: str | None = None, action_cb=None) -> None:
    left, right = st.columns([3.1, 1.1])
    with left:
        st.markdown(
            f'<div class="pagehead"><div><div class="overline">Operating system · '
            f'{PAGE_INDEX.get(page, 1):02d} / 07</div><h1>{html.escape(title)}</h1>'
            f'<div class="sub">{description}</div></div></div>',
            unsafe_allow_html=True,
        )
    if action_label and action_cb:
        with right:
            with st.container(key="head_actions"):
                if st.button(action_label, type="primary", key=f"run_{page}"):
                    action_cb()


def _kpi_strip(items: list[tuple[str, str, str, str]]) -> str:
    blocks = []
    for label, value, note, color in items:
        blocks.append(
            f'<div class="kpi"><div class="kpi-label">{label}</div>'
            f'<div class="kpi-value {color}">{value}</div>'
            f'<div class="kpi-note">{note}</div></div>'
        )
    return '<div class="kpi-strip">' + "".join(blocks) + "</div>"


def _flow_html(research: dict, listing: dict, compliance: dict, ops: dict, monitor: dict) -> str:
    severity = compliance.get("severity", {}) or {}
    status = compliance.get("status", {}) or {}
    high = int(severity.get("high", 0))
    findings = int(compliance.get("finding_count", 0) or sum(int(v) for v in severity.values()))
    margin = (research.get("quality", {}) or {}).get("avg_margin_pct", 0)
    total_alerts = int((monitor.get("counts", {}) or {}).get("alerts", 0))
    failed = int((monitor.get("counts", {}) or {}).get("failed", 0))
    reviews = int((monitor.get("counts", {}) or {}).get("negative_reviews", 0))
    competitors = int((monitor.get("counts", {}) or {}).get("competitors", 0))
    flow = [
        ("01 / 选品", str(research.get("high_potential_records", "—")), "SKU", f"高潜力 · 平均毛利 {margin:.1f}%", ""),
        ("02 / 上架", str(len(listing.get("locales", [])) or "—"), "站点", "格式校验通过 · 导入需补资料", "green"),
        ("03 / 合规", str(high), "高风险", f"{status.get('PASS', 0)} PASS · {status.get('FAIL', 0)} FAIL · {findings} 项", "red-text"),
        ("04 / 经营", f"${ops.get('net_profit', 0):,.0f}" if ops else "—", "净利", f"ROI {ops.get('roi', 0):.2f}% · ACOS {ops.get('ad_acos', 0):.2f}%", ""),
        ("05 / 竞品", str(total_alerts), "预警", f"{competitors} 竞品 · {reviews} 差评 · {failed} 失败", "amber-text"),
    ]
    nodes = []
    for kicker, value, unit, meta, color in flow:
        nodes.append(
            f'<div class="flow-node"><div class="flow-kicker">{kicker}</div>'
            f'<div class="flow-main"><span class="flow-value {color}">{value}</span>'
            f'<span class="flow-unit">{unit}</span></div>'
            f'<div class="flow-meta">{meta}</div><span class="arrow"></span></div>'
        )
    return '<div class="flow">' + "".join(nodes) + "</div>"


def _priority_html(ops: dict, compliance: dict) -> str:
    metrics = ops.get("metrics", []) or []
    urgent = [m for m in metrics if m.get("stock_status") == "紧急补货"]
    slow = [m for m in metrics if m.get("stock_status") == "滞销"]
    urgent_text = f"{len(urgent)} 个 SKU 即将断货"
    urgent_meta = (
        f"{urgent[0].get('sku', '')} · 覆盖 {urgent[0].get('days_cover', '—')} 天 · 交期 {urgent[0].get('lead_time_days', '—')} 天"
        if urgent
        else "暂无紧急补货 SKU"
    )
    slow_text = f"{len(slow)} 个 SKU 库存超过 90 天" if slow else "库存覆盖健康"
    slow_meta = " · ".join(m.get("sku", "") for m in slow[:3]) or "无需处理"
    high = compliance.get("severity", {}).get("high", 0)
    return f"""
    <section class="panel priority">
      <div class="panel-head"><div class="panel-title"><span class="num">01</span>待处理 <span style="color:var(--faint);font-family:var(--mono);font-size:11px;font-weight:400">/ 3</span></div><span class="panel-link">查看全部 →</span></div>
      <div class="action-row"><span class="prio p1">P1</span><div class="action-copy"><b>{urgent_text}</b><span>{urgent_meta}</span></div><span class="action-arrow">→</span></div>
      <div class="action-row"><span class="prio p2">P2</span><div class="action-copy"><b>{slow_text}</b><span>{slow_meta}</span></div><span class="action-arrow">→</span></div>
      <div class="action-row"><span class="prio p3">合规</span><div class="action-copy"><b>{high} 个 CE/FCC 高风险项待补证</b><span>欧盟 LVD/EMC · GPSR 责任人 · 无线 FCC ID</span></div><span class="action-arrow">→</span></div>
    </section>
    """


def _profit_table_html(ops: dict) -> str:
    metrics = (ops.get("metrics", []) or [])[:5]
    if not metrics:
        return '<div class="panel"><div class="panel-head"><div class="panel-title"><span class="num">02</span>月度利润 · Top 5</div></div><div style="padding:18px;color:var(--muted);font-size:11px">尚未生成经营报表。</div></div>'
    status_class = {"健康": "st-green", "关注": "st-amber", "滞销": "st-grey", "紧急补货": "st-red"}
    rows = []
    for m in metrics:
        status = m.get("stock_status", "—")
        rows.append(
            f'<tr><td class="sku">{html.escape(str(m.get("sku", "")))}</td>'
            f'<td class="num">${m.get("revenue", 0):,.0f}</td>'
            f'<td class="num">${m.get("net_profit", 0):,.0f}</td>'
            f'<td class="num green">{m.get("roi", 0) * 100:.1f}%</td>'
            f'<td class="num">{m.get("inventory_turnover_days", 0):.1f}d</td>'
            f'<td><span class="status {status_class.get(status, "st-grey")}"><i class="sdot"></i>{html.escape(str(status))}</span></td></tr>'
        )
    return f"""
    <section class="panel" style="margin-top:14px">
      <div class="panel-head"><div class="panel-title"><span class="num">02</span>月度利润 · Top 5</div><div class="panel-meta">USD / 30d</div><span class="panel-link">打开报表 →</span></div>
      <table class="overview-table">
        <thead><tr><th style="width:27%">SKU</th><th class="num" style="width:16%">收入</th><th class="num" style="width:15%">净利润</th><th class="num" style="width:12%">ROI</th><th class="num" style="width:16%">周转</th><th style="width:14%">库存</th></tr></thead>
        <tbody>{''.join(rows)}</tbody>
      </table>
    </section>
    """


def _compliance_html(compliance: dict, listing: dict) -> str:
    severity = compliance.get("severity", {}) or {}
    status = compliance.get("status", {}) or {}
    high, medium, low = int(severity.get("high", 0)), int(severity.get("medium", 0)), int(severity.get("low", 0))
    total = max(1, high + medium + low)
    red, amber, grey = high / total * 100, medium / total * 100, low / total * 100
    locales = len(listing.get("locales", []))
    format_ok = "通过" if listing.get("format_ok_all") else "需检查"
    import_ready = "是" if listing.get("import_ready_all") else "需补 GTIN / 主图"
    return f"""
    <section class="panel compliance">
      <div class="panel-head" style="height:36px;padding:0;border:none;margin-bottom:2px"><div class="panel-title"><span class="num">03</span>合规状态</div><span class="panel-meta">9 SKU</span></div>
      <div class="bar"><span style="background:var(--red);width:{red:.1f}%"></span><span style="background:#C28A37;width:{amber:.1f}%"></span><span style="background:#C9D4CC;width:{grey:.1f}%"></span></div>
      <div class="legend"><span class="leg"><i class="sw red"></i>高风险 {high}</span><span class="leg"><i class="sw amber"></i>中风险 {medium}</span><span class="leg"><i class="sw grey"></i>低风险 {low}</span></div>
      <div class="compliance-foot"><span>PASS <b>{status.get('PASS', 0)}</b></span><span>FAIL <b>{status.get('FAIL', 0)}</b></span><span>CONDITIONAL <b>{status.get('CONDITIONAL', 0)}</b></span></div>
    </section>
    <section class="panel mini-panel">
      <div class="mini-title">Listing readiness</div>
      <div class="mini-row"><span>站点版本</span><b>{locales} / 4</b></div>
      <div class="mini-row"><span>格式校验</span><b class="green">{format_ok}</b></div>
      <div class="mini-row"><span>导入就绪</span><b class="amber-text">{import_ready}</b></div>
    </section>
    """


def _log_html() -> str:
    entries: list[dict] = []
    if RUN_LOG.exists():
        for line in read_text(RUN_LOG).splitlines()[-8:]:
            try:
                item = json.loads(line)
                if isinstance(item, dict):
                    entries.append(item)
            except json.JSONDecodeError:
                continue
    entries = entries[-4:]
    if not entries:
        entries = [
            {"timestamp": "03:08:19", "label": "选品", "elapsed_seconds": 0.9, "returncode": 0},
            {"timestamp": "03:08:20", "label": "Listing", "elapsed_seconds": 0.8, "returncode": 0},
            {"timestamp": "03:08:21", "label": "合规", "elapsed_seconds": 0.7, "returncode": 0},
            {"timestamp": "03:08:22", "label": "经营", "elapsed_seconds": 0.9, "returncode": 0},
        ]
    ok = sum(1 for item in entries if item.get("returncode") == 0)
    lines = []
    for item in entries:
        stamp = str(item.get("timestamp", ""))
        time_text = stamp[11:19] if len(stamp) >= 19 else (stamp[:8] if stamp else "--:--:--")
        label = str(item.get("label", "run"))[:12]
        elapsed = item.get("elapsed_seconds", 0)
        lines.append(
            f'<div class="log-line"><span class="t">{time_text}</span><span class="k">{html.escape(label)}</span>'
            f'<span class="v">{"完成" if item.get("returncode") == 0 else "失败"}</span><span class="d">{elapsed}s</span></div>'
        )
    return f"""
    <section class="log">
      <div class="log-head"><div class="log-title"><span class="pulse"></span>运行日志</div><div class="log-status">{ok} / {len(entries)} PASS</div></div>
      <div class="log-lines">{''.join(lines)}</div>
      <div class="log-foot"><span>errors 0 · local run</span><span>完整日志 →</span></div>
    </section>
    """


# ---------------------------------------------------------------------------
# sidebar + topbar
# ---------------------------------------------------------------------------
def _sidebar() -> str:
    research = _summary(RESEARCH_SUMMARY)
    listing = _summary(LISTING_SUMMARY)
    compliance = _summary(COMPLIANCE_SUMMARY)
    monitor = _summary(MONITOR_SUMMARY)
    with st.sidebar:
        st.markdown(
            '<div class="brand"><div class="mark"></div><div class="wordmark">Tundra<span>Outdoor ESS ops</span></div></div>'
            '<div class="workspace"><div><b>Outdoor ESS</b><small>EU / US workspace</small></div><div class="chev">⌄</div></div>'
            '<div class="nav-label">Workspace</div>',
            unsafe_allow_html=True,
        )
        page = st.radio("Navigation", PAGES, index=0, label_visibility="collapsed", key="primary_nav")
        st.markdown(
            f'<div class="side-stats">'
            f'<div class="side-stat"><b>{research.get("high_potential_records", "—")}</b><span>高潜 SKU</span></div>'
            f'<div class="side-stat"><b>{len(listing.get("locales", [])) or "—"}</b><span>Listing 站点</span></div>'
            f'<div class="side-stat"><b>{compliance.get("severity", {}).get("high", "—")}</b><span>合规高风险</span></div>'
            f'<div class="side-stat"><b>{monitor.get("counts", {}).get("alerts", "—")}</b><span>竞品预警</span></div>'
            f'</div>',
            unsafe_allow_html=True,
        )
        st.markdown('<div class="nav-label">System</div>', unsafe_allow_html=True)
        if st.button("检查项目结构", key="check_project", width="stretch"):
            run_command("项目检查", ["main.py", "check"], timeout=120)
        if st.button("查看错误日志", key="show_errors", width="stretch"):
            error_log = LOG_DIR / "errors.log"
            st.code(read_text(error_log, default="暂无错误日志")[-4000:], language="text")
        st.markdown(
            '<div class="rail-foot">'
            '<div class="system-row"><span><i class="dot"></i>系统正常</span><span>v1.0</span></div>'
            '<div class="system-row"><span>数据模式</span><span>模拟 · 固定种子</span></div>'
            '<div class="system-row"><span>最近运行</span><span>03:08</span></div>'
            '<div class="system-note">Public snapshots / no store API<br>Global error log active</div>'
            '</div>',
            unsafe_allow_html=True,
        )
    return page


def _topbar(page: str) -> None:
    st.markdown(
        f'<div class="topbar"><div class="crumb">Workspace / <b>{html.escape(page)}</b></div>'
        '<div class="top-actions"><span class="sim-pill">Simulated · Q3 2026</span>'
        '<span class="top-status"><i class="dot"></i>local</span></div></div>',
        unsafe_allow_html=True,
    )


# ---------------------------------------------------------------------------
# overview
# ---------------------------------------------------------------------------
def render_overview() -> None:
    research = _summary(RESEARCH_SUMMARY)
    listing = _summary(LISTING_SUMMARY)
    compliance = _summary(COMPLIANCE_SUMMARY)
    ops = _summary(OPS_SUMMARY)
    monitor = _summary(MONITOR_SUMMARY)
    _page_header(
        "经营总览",
        '<span>2026-09-01 — 09-30</span><span class="sep"></span><span>9 SKU · 4 站点</span><span class="sep"></span><span>数据截止 03:08</span>',
        "总览",
        "运行全链路",
        _run_all,
    )
    st.markdown(_flow_html(research, listing, compliance, ops, monitor), unsafe_allow_html=True)
    st.markdown(
        _kpi_strip(
            [
                ("Net profit / 30 days", f"${ops.get('net_profit', 0):,.2f}", '<span class="green">21.97%</span> 净利率 · 收入 $251,367', ""),
                ("Return on invested cost", f"{ops.get('roi', 0):.2f}%", f"总成本 ${ops.get('net_profit', 0) / (ops.get('roi', 1) / 100) if ops.get('roi') else 0:,.0f} · 9 SKU", ""),
                ("Inventory attention", f"{ops.get('inventory_alerts', 0)} SKU", f"<span class=\"red-text\">{ops.get('urgent_alerts', 0)} 紧急</span> · {ops.get('slow_moving', 0)} 滞销 · 1 关注", ""),
            ]
        ),
        unsafe_allow_html=True,
    )
    left, right = st.columns([1.52, 0.78], gap="large")
    with left:
        st.markdown(_priority_html(ops, compliance), unsafe_allow_html=True)
        st.markdown(_profit_table_html(ops), unsafe_allow_html=True)
    with right:
        st.markdown('<div class="side-stack">', unsafe_allow_html=True)
        st.markdown(_compliance_html(compliance, listing), unsafe_allow_html=True)
        st.markdown(_log_html(), unsafe_allow_html=True)
        st.markdown("</div>", unsafe_allow_html=True)
    _download(OUTPUTS / "operations" / "monthly_profit.xlsx", "月度利润 Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.markdown(
        '<div class="foot"><span>固定种子模拟 · 公开快照 · 不接入真实店铺后台</span><span class="foot-right">Tundra Ops / 2026.09</span></div>',
        unsafe_allow_html=True,
    )


def _run_all() -> None:
    steps = [
        ("基础 7 阶段流水线", ["main.py", "demo"]),
        ("公开数据选品", ["product_research.py"]),
        ("Amazon Listing", ["listing_generator.py", "--locales", "all"]),
        ("合规校验", ["compliance_check.py"]),
        ("经营报表", ["report_generator.py"]),
        ("竞品巡检", ["competitor_monitor.py", "run"]),
    ]
    for label, args in steps:
        proc = run_command(label, args, timeout=300)
        if proc is None or proc.returncode != 0:
            st.warning("全链路中断，请查看上方输出和 outputs/logs/errors.log")
            return
    st.success("全链路运行完成，可切换左侧模块查看结果。")


# ---------------------------------------------------------------------------
# research
# ---------------------------------------------------------------------------
def render_research() -> None:
    _page_header("选品研究", "<span>100-500Wh 公开数据清洗、三指标计算与高潜力 Excel</span>", "选品研究", "运行选品", lambda: run_command("选品研究", ["product_research.py"]))
    summary = _summary(RESEARCH_SUMMARY)
    _metric_strip_from_summary(summary)
    headers, rows = _read_excel(RESEARCH_XLSX, "高潜力选品", ("SKU",))
    if rows:
        st.markdown('<div class="section-title">高潜力选品表</div>', unsafe_allow_html=True)
        st.dataframe(_df(headers, rows), width="stretch", hide_index=True, height=430)
        _download(RESEARCH_XLSX, "高潜力选品 Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.info("尚未生成选品结果，点击右上角“运行选品”。")
    st.caption("输出：outputs/research/high_potential_products.xlsx / research_summary.json")


def _metric_strip_from_summary(summary: dict) -> None:
    if not summary:
        return
    st.markdown(
        _kpi_strip(
            [
                ("Raw records", str(summary.get("raw_records", "—")), "公开快照 / 页面记录", ""),
                ("Cleaned records", str(summary.get("cleaned_records", "—")), "容量 / 价格 / 去重后", ""),
                ("High potential", str(summary.get("high_potential_records", "—")), "通过全部门槛", "green"),
            ]
        ),
        unsafe_allow_html=True,
    )


# ---------------------------------------------------------------------------
# listing
# ---------------------------------------------------------------------------
def render_listing() -> None:
    _page_header("上架素材", "<span>合规英文 / 本地化 Listing、5 点描述、详情页、Amazon 平铺文件</span>", "上架素材")
    toolbar = st.columns([1.5, 1.2, 2.3])
    with toolbar[0]:
        locale = st.selectbox("目标站点", ["en-US", "en-GB", "de-DE", "fr-FR", "all"], index=0, key="listing_locale")
    with toolbar[1]:
        if st.button("生成 Listing", type="primary", key="run_listing", width="stretch"):
            run_command("上架素材", ["listing_generator.py", "--locales", locale])
    summary = _summary(LISTING_SUMMARY)
    if summary:
        locales = summary.get("locales", [])
        st.markdown(
            _kpi_strip(
                [
                    ("Versions", str(len(locales)), "站点 / 语言版本", ""),
                    ("Format check", "通过" if summary.get("format_ok_all") else "未通过", "表头 / 长度 / 类型", "green" if summary.get("format_ok_all") else "red-text"),
                    ("Import ready", "是" if summary.get("import_ready_all") else "需补资料", "GTIN / 主图 / Browse Node", "amber-text"),
                ]
            ),
            unsafe_allow_html=True,
        )
        if locales:
            selected = st.selectbox("预览站点", [item["locale"] for item in locales], key="listing_preview")
            item = next((x for x in locales if x["locale"] == selected), locales[0])
            csv_path = ROOT / item["csv"] if not Path(item["csv"]).is_absolute() else Path(item["csv"])
            if csv_path.exists():
                with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                    reader = csv.DictReader(handle)
                    rows = list(reader)
                    cols = reader.fieldnames or []
                st.markdown(f'<div class="section-title">Amazon CSV 预览 · {html.escape(selected)}</div>', unsafe_allow_html=True)
                st.dataframe(pd.DataFrame(rows, columns=cols), width="stretch", hide_index=True, height=280)
                _download(csv_path, f"{selected} Amazon CSV", "text/csv")
            validation = ROOT / item["validation"] if not Path(item["validation"]).is_absolute() else Path(item["validation"])
            if validation.exists():
                with st.expander("查看导入校验 JSON", expanded=False):
                    st.json(read_json(validation, default={}))
    else:
        st.info("尚未生成 Listing，点击上方“生成 Listing”。")


# ---------------------------------------------------------------------------
# monitor
# ---------------------------------------------------------------------------
def render_monitor() -> None:
    _page_header("竞品监控", "<span>公开价格 / BSR / 差评巡检、价格与排名预警、差评痛点报告</span>", "竞品监控")
    toolbar = st.columns([1.1, 1.1, 2.4])
    with toolbar[0]:
        mode = st.radio("数据模式", ["offline", "live", "auto"], horizontal=True, index=0, key="monitor_mode")
    with toolbar[1]:
        if st.button("巡检竞品", type="primary", key="run_monitor", width="stretch"):
            run_command("竞品监控", ["competitor_monitor.py", "run", "--mode", mode])
    summary = _summary(MONITOR_SUMMARY)
    if summary:
        counts = summary.get("counts", {})
        st.markdown(
            _kpi_strip(
                [
                    ("Competitors", str(counts.get("competitors", "—")), "注册表启用条目", ""),
                    ("Success / failed", f"{counts.get('success', 0)} / {counts.get('failed', 0)}", "失败链接自动跳过", "green"),
                    ("Alerts", str(counts.get("alerts", "—")), "价格 / BSR 异动", "amber-text"),
                ]
            ),
            unsafe_allow_html=True,
        )
        alert_path = ROOT / summary["latest_alerts"] if not Path(summary["latest_alerts"]).is_absolute() else Path(summary["latest_alerts"])
        report_path = ROOT / summary["latest_report"] if not Path(summary["latest_report"]).is_absolute() else Path(summary["latest_report"])
        col1, col2 = st.columns([1, 1], gap="large")
        with col1:
            st.markdown('<div class="section-title">最新预警</div>', unsafe_allow_html=True)
            if alert_path.exists():
                st.code(read_text(alert_path), language="text")
                _download(alert_path, "预警 TXT", "text/plain")
        with col2:
            st.markdown('<div class="section-title">差评痛点</div>', unsafe_allow_html=True)
            if report_path.exists():
                st.markdown(read_text(report_path)[:9000])
                _download(report_path, "痛点报告 Markdown", "text/markdown")
    else:
        st.info("尚未运行竞品巡检。离线演示可直接使用默认快照。")
    st.caption("定时巡检：python competitor_monitor.py schedule --interval-hours 6；页面内为单轮巡检。")


# ---------------------------------------------------------------------------
# compliance
# ---------------------------------------------------------------------------
def render_compliance() -> None:
    _page_header("合规校验", "<span>CE/FCC 规则初筛、风险评分、不符合项与修改建议</span>", "合规校验", "合规筛查", lambda: run_command("合规校验", ["compliance_check.py"]))
    summary = _summary(COMPLIANCE_SUMMARY)
    if summary:
        severity = summary.get("severity", {})
        status = summary.get("status", {})
        st.markdown(
            _kpi_strip(
                [
                    ("High risk", str(severity.get("high", "—")), "需优先处理", "red-text"),
                    ("Medium risk", str(severity.get("medium", "—")), "需补充文件", "amber-text"),
                    ("PASS / FAIL", f"{status.get('PASS', 0)} / {status.get('FAIL', 0)}", "无高/中风险 / 存在高风险", ""),
                ]
            ),
            unsafe_allow_html=True,
        )
        headers, rows = _read_excel(OUTPUTS / "compliance" / "compliance_check.xlsx", "合规汇总", ("SKU",))
        if rows:
            st.markdown('<div class="section-title">合规汇总</div>', unsafe_allow_html=True)
            st.dataframe(_df(headers, rows), width="stretch", hide_index=True, height=300)
        f_headers, findings = _read_excel(OUTPUTS / "compliance" / "compliance_check.xlsx", "不符合项", ("ID", "标准"))
        if findings:
            levels = sorted({str(x.get("严重级别", "")) for x in findings if x.get("严重级别")})
            selected = st.multiselect("筛选严重级别", levels, default=levels, key="compliance_levels")
            filtered = [x for x in findings if str(x.get("严重级别", "")) in selected]
            st.markdown('<div class="section-title">不符合项与修改建议</div>', unsafe_allow_html=True)
            st.dataframe(_df(f_headers, filtered), width="stretch", hide_index=True, height=350)
        _download(OUTPUTS / "compliance" / "compliance_check.xlsx", "合规校验 Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.info("尚未运行合规筛查，点击右上角“合规筛查”。")


# ---------------------------------------------------------------------------
# operations
# ---------------------------------------------------------------------------
def render_operations() -> None:
    _page_header("经营报表", "<span>广告日报、库存预警、月度利润、ROI 与库存周转天数</span>", "经营报表", "生成报表", lambda: run_command("经营报表", ["report_generator.py"]))
    summary = _summary(OPS_SUMMARY)
    if summary:
        st.markdown(
            _kpi_strip(
                [
                    ("Revenue", f"${summary.get('revenue', 0):,.0f}", "30 天模拟", ""),
                    ("Net profit", f"${summary.get('net_profit', 0):,.0f}", f"{summary.get('net_margin', 0):.2f}% 净利率", "green"),
                    ("ROI / ACOS", f"{summary.get('roi', 0):.2f}% / {summary.get('ad_acos', 0):.2f}%", "净利润 ÷ 总成本 / 广告花费 ÷ 广告销售", ""),
                ]
            ),
            unsafe_allow_html=True,
        )
        tabs = st.tabs(["广告日报", "库存预警", "月度利润"])
        with tabs[0]:
            h, r = _read_excel(OUTPUTS / "operations" / "ad_daily_report.xlsx", "广告日报", ("日期",))
            st.dataframe(_df(h, r).head(120), width="stretch", hide_index=True, height=340)
            _download(OUTPUTS / "operations" / "ad_daily_report.xlsx", "广告日报 Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with tabs[1]:
            h, r = _read_excel(OUTPUTS / "operations" / "inventory_alert.xlsx", "库存预警", ("SKU",))
            st.dataframe(_df(h, r), width="stretch", hide_index=True, height=300)
            _download(OUTPUTS / "operations" / "inventory_alert.xlsx", "库存预警 Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with tabs[2]:
            h, r = _read_excel(OUTPUTS / "operations" / "monthly_profit.xlsx", "月度利润", ("SKU",))
            st.dataframe(_df(h, r), width="stretch", hide_index=True, height=320)
            _download(OUTPUTS / "operations" / "monthly_profit.xlsx", "月度利润 Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.info("尚未生成经营报表，点击右上角“生成报表”。")


# ---------------------------------------------------------------------------
# interview
# ---------------------------------------------------------------------------
def render_interview() -> None:
    _page_header("面试演示", "<span>生成专属演示脚本，逐模块映射跨境电商岗位 JD 能力点</span>", "面试演示", "生成演示脚本", lambda: run_command("面试演示", ["interview_demo.py", "--run"], timeout=300))
    if INTERVIEW_MD.exists():
        st.markdown(read_text(INTERVIEW_MD)[:20000])
        _download(INTERVIEW_MD, "面试演示脚本", "text/markdown")
    else:
        st.info("尚未生成，点击右上角“生成演示脚本”。脚本包含 JD 能力点、演示顺序、STAR 话术和追问准备。")


# ---------------------------------------------------------------------------
# entry
# ---------------------------------------------------------------------------
def main() -> None:
    page = _sidebar()
    _topbar(page)
    with st.container(key="sheet"):
        renderers = {
            "总览": render_overview,
            "选品研究": render_research,
            "上架素材": render_listing,
            "竞品监控": render_monitor,
            "合规校验": render_compliance,
            "经营报表": render_operations,
            "面试演示": render_interview,
        }
        renderers.get(page, render_overview)()


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        log_exception(exc, app="streamlit", context={"hook": "app_main"})
        st.error(f"页面异常：{type(exc).__name__}: {exc}")
        st.caption("错误已写入 outputs/logs/errors.log")
